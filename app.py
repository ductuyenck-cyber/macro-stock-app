# StockAI Độc lập 1.2 — all application code in one file.
from datetime import datetime
from pathlib import Path
from copy import deepcopy
import re
import pandas as pd
import streamlit as st
"""Pure calculation engine. Missing input is None, never an implicit zero."""
import math

FIELDS = {
 'sales': ('Doanh thu thuần', 'Báo cáo KQKD: doanh thu thuần, không dùng tổng doanh thu.'),
 'receivables': ('Phải thu khách hàng thuần', 'Chỉ phải thu khách hàng; dùng nhất quán sau dự phòng cho hai năm.'),
 'cogs': ('Giá vốn hàng bán', 'Nhập số dương; không nhập dấu âm theo cách trình bày của BCTC.'),
 'current_assets': ('Tài sản ngắn hạn', 'Bảng cân đối kế toán: tổng tài sản ngắn hạn.'),
 'ppe': ('TSCĐ hữu hình còn lại', 'Nguyên giá trừ hao mòn lũy kế; không cộng BĐS đầu tư, XDCB dở dang, đầu tư tài chính.'),
 'depreciation': ('Khấu hao TSCĐ hữu hình trong năm', 'Đọc thuyết minh, khớp phạm vi PP&E. Không dùng hao mòn lũy kế.'),
 'sga': ('Chi phí bán hàng + quản lý', 'Cộng chi phí bán hàng và chi phí quản lý doanh nghiệp, nhập số dương.'),
 'long_liabilities': ('Tổng nợ phải trả dài hạn', 'Dùng cho LVGI của M-Score, không chỉ vay dài hạn.'),
 'current_liabilities': ('Tổng nợ ngắn hạn', 'Bảng cân đối kế toán: tổng nợ ngắn hạn.'),
 'income': ('LNST dùng cho F-Score', 'Dùng nhất quán lợi nhuận toàn doanh nghiệp. Mô hình gốc dùng lợi nhuận trước khoản bất thường.'),
 'continuing_income': ('Lợi nhuận hoạt động tiếp tục sau thuế', 'Dùng cho TATA. Nếu dùng LNST thay thế, phải đánh dấu xấp xỉ ở thông tin báo cáo.'),
 'cfo': ('Lưu chuyển tiền thuần từ HĐKD', 'Lưu chuyển tiền tệ: dòng tiền thuần từ hoạt động kinh doanh; có thể âm.'),
 'assets': ('Tổng tài sản cuối năm', 'Cần thêm giá trị cuối t-2 để tính ROA/vòng quay đầu kỳ.'),
 'long_debt': ('Vay dài hạn gồm phần đến hạn trả', 'Dùng cho F-Score. Gồm phần vay dài hạn tái phân loại vào nợ ngắn hạn; không dùng toàn bộ nợ phải trả dài hạn.'),
}
YEARS = ('t','t1','t2')
WEIGHTS = {'DSRI': .920, 'GMI': .528, 'AQI': .404, 'SGI': .892, 'DEPI': .115, 'SGAI': -.172, 'LVGI': -.327, 'TATA': 4.679}

def blank():
 return {'meta': {'ticker':'', 'year':2025, 'unit':'Triệu VND', 'scope':'Hợp nhất', 'sector':'Doanh nghiệp phi tài chính', 'source':'', 'issuance':'Chưa xác định', 'f_income_proxy':True, 'm_income_proxy':True, 'demo':False}, 'periods':{y:{k:None for k in FIELDS} for y in YEARS}}

def demo():
 d=blank();d['meta'].update(ticker='DEMO',source='Dữ liệu giả lập để học; không phải số liệu doanh nghiệp.',issuance='Không',demo=True)
 for y in ['t','t1']:
  d['periods'][y].update(sales=100.,receivables=10.,cogs=60.,current_assets=40.,ppe=30.,depreciation=5.,sga=10.,long_liabilities=20.,current_liabilities=20.,income=10.,continuing_income=10.,cfo=12.,assets=100.,long_debt=20.)
 d['periods']['t2']['assets']=100.
 return d

def calculate(data):
 errors=[]; warnings=[]; p=data['periods']; meta=data['meta']
 for y in YEARS:
  for k in FIELDS:
   v=p[y].get(k)
   if v is not None and (isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v)):
    errors.append(f'{y}: {FIELDS[k][0]} phải là số hữu hạn.'); continue
   if v is not None and k not in ('income','continuing_income','cfo') and v<0:
    errors.append(f'{y}: {FIELDS[k][0]} không được âm.')
  for k in ['assets','sales']:
   if p[y].get(k) is not None and p[y][k]<=0: errors.append(f'{y}: {FIELDS[k][0]} phải lớn hơn 0.')
  ca,pp,ta=(p[y].get(k) for k in ['current_assets','ppe','assets'])
  if all(v is not None for v in [ca,pp,ta]) and ca+pp>ta+max(1,abs(ta))*1e-9: errors.append(f'{y}: tài sản ngắn hạn + PP&E vượt tổng tài sản.')
  rec=p[y].get('receivables')
  if rec is not None and ca is not None and rec>ca: warnings.append(f'{y}: phải thu vượt tài sản ngắn hạn; kiểm tra có cộng phải thu dài hạn không.')
  ll,cl=p[y].get('long_liabilities'),p[y].get('current_liabilities')
  if all(v is not None for v in [ll,cl,ta]) and ll+cl>ta: warnings.append(f'{y}: nợ vượt tài sản, có thể vốn chủ âm; cần kiểm tra.')
 if meta.get('demo'): warnings.append('Đang dùng dữ liệu minh họa, không phải số liệu cổ phiếu thực.')
 if not str(meta.get('source','')).strip(): warnings.append('Chưa ghi nguồn BCTC và trang/thuyết minh.')
 if meta.get('f_income_proxy'): warnings.append('F-Score dùng LNST thay lợi nhuận trước khoản bất thường: bản điều chỉnh cần đối chiếu thuyết minh.')
 if meta.get('m_income_proxy'): warnings.append('TATA dùng LNST thay lợi nhuận hoạt động tiếp tục: M-Score xấp xỉ.')
 out={'errors':errors,'warnings':warnings,'f_rows':[],'m_rows':[], 'f_score':None,'m_score':None}
 if errors:return out
 if meta.get('sector') != 'Doanh nghiệp phi tài chính':
  errors.append('Mô hình này không chấm ngân hàng, bảo hiểm, công ty chứng khoán.');return out
 def get(y,k):
  v=p[y].get(k)
  if v is None:raise ValueError(f'Thiếu {FIELDS[k][0]} ({y}).')
  return v
 def div(a,b,label):
  if b==0:raise ValueError(f'{label}: mẫu số bằng 0.')
  return a/b
 def ratio(y,n,d):return div(get(y,n),get(y,d),FIELDS[n][0]+'/'+FIELDS[d][0])
 def margin(y):return div(get(y,'sales')-get(y,'cogs'),get(y,'sales'),'Biên lợi nhuận gộp')
 def avg(y,prev):return (get(y,'assets')+get(prev,'assets'))/2
 def issuance():
  val=meta.get('issuance')
  if val not in ['Có','Không']:raise ValueError('Chưa xác định phát hành cổ phần huy động vốn.')
  return (1 if val=='Không' else 0,1,'Không = đạt; Có = không đạt')
 # Each criterion returns observed current/prior value; no missing -> false conversion.
 criteria=[
  ('LNST dương',lambda:(get('t','income'),0,'LNST > 0'),lambda a,b:a>b),
  ('CFO dương',lambda:(get('t','cfo'),0,'CFO > 0'),lambda a,b:a>b),
  ('ROA tăng',lambda:(div(get('t','income'),get('t1','assets'),'ROA t'),div(get('t1','income'),get('t2','assets'),'ROA t-1'),'LN / tài sản đầu kỳ; t > t-1'),lambda a,b:a>b),
  ('CFO lớn hơn LNST',lambda:(get('t','cfo'),get('t','income'),'CFO > LNST'),lambda a,b:a>b),
  ('Đòn bẩy giảm',lambda:(div(get('t','long_debt'),avg('t','t1'),'Đòn bẩy t'),div(get('t1','long_debt'),avg('t1','t2'),'Đòn bẩy t-1'),'Vay dài hạn / tài sản bình quân; t < t-1'),lambda a,b:a<b),
  ('Thanh toán hiện hành tăng',lambda:(ratio('t','current_assets','current_liabilities'),ratio('t1','current_assets','current_liabilities'),'CA / CL; t > t-1'),lambda a,b:a>b),
  ('Không phát hành huy động vốn',issuance,lambda a,b:a==b),
  ('Biên lợi nhuận gộp tăng',lambda:(margin('t'),margin('t1'),'(Doanh thu − giá vốn)/doanh thu; t > t-1'),lambda a,b:a>b),
  ('Vòng quay tài sản tăng',lambda:(div(get('t','sales'),get('t1','assets'),'Vòng quay t'),div(get('t1','sales'),get('t2','assets'),'Vòng quay t-1'),'Doanh thu / tài sản đầu kỳ; t > t-1'),lambda a,b:a>b)]
 for name,fn,predicate in criteria:
  try:
   a,b,rule=fn(); point=int(predicate(a,b)); row={'Tiêu chí':name,'Điểm':point,'Giá trị t':a,'Đối chiếu':b,'Quy tắc / Lý do':rule}
  except ValueError as exc:row={'Tiêu chí':name,'Điểm':None,'Giá trị t':None,'Đối chiếu':None,'Quy tắc / Lý do':str(exc)}
  out['f_rows'].append(row)
 if all(r['Điểm'] is not None for r in out['f_rows']):out['f_score']=sum(r['Điểm'] for r in out['f_rows'])
 def rel(n,d,label):return div(ratio('t',n,d),ratio('t1',n,d),label)
 def gmi():
  a,b=margin('t1'),margin('t')
  if min(a,b)<=0:raise ValueError('Biên lợi nhuận gộp không dương; không diễn giải GMI theo mô hình thông thường.')
  return div(a,b,'GMI')
 def quality(y):return 1-div(get(y,'current_assets')+get(y,'ppe'),get(y,'assets'),'AQI')
 def dep(y):return div(get(y,'depreciation'),get(y,'ppe')+get(y,'depreciation'),'DEPI')
 def lev(y):return div(get(y,'long_liabilities')+get(y,'current_liabilities'),get(y,'assets'),'LVGI')
 funcs={
  'DSRI':lambda:rel('receivables','sales','DSRI'), 'GMI':gmi,
  'AQI':lambda:div(quality('t'),quality('t1'),'AQI'),
  'SGI':lambda:div(get('t','sales'),get('t1','sales'),'SGI'),
  'DEPI':lambda:div(dep('t1'),dep('t'),'DEPI'),
  'SGAI':lambda:rel('sga','sales','SGAI'),
  'LVGI':lambda:div(lev('t'),lev('t1'),'LVGI'),
  'TATA':lambda:div(get('t','continuing_income')-get('t','cfo'),get('t','assets'),'TATA')}
 for name,fn in funcs.items():
  try:
   val=fn()
   if not math.isfinite(val):raise ValueError('Kết quả không hữu hạn.')
   row={'Biến':name,'Giá trị':val,'Hệ số':WEIGHTS[name],'Đóng góp':val*WEIGHTS[name],'Lý do':''}
  except ValueError as exc:row={'Biến':name,'Giá trị':None,'Hệ số':WEIGHTS[name],'Đóng góp':None,'Lý do':str(exc)}
  out['m_rows'].append(row)
 if all(r['Giá trị'] is not None for r in out['m_rows']):out['m_score']=-4.84+sum(r['Đóng góp'] for r in out['m_rows'])
 return out

"""Workbook interchange used by the application at runtime."""
from io import BytesIO
import math
from zipfile import ZipFile, BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


BOOL_KEYS={'demo','f_income_proxy','m_income_proxy'}
def safe_text(v):
 return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v

def format_sheet(ws):
 ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='123047')
 for col in ws.columns:
  letter=col[0].column_letter;ws.column_dimensions[letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col)+2))
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(cell.value,(int,float)):cell.number_format='#,##0.000;[Red](#,##0.000);0'
  ws.row_dimensions[row[0].row].height=32

def export_workbook(data, result=None):
 wb=Workbook();ws=wb.active;ws.title='Du_lieu'
 ws.append(['field','Chỉ tiêu','t','t1','t2','Hướng dẫn'])
 for key,(label,help_) in FIELDS.items():ws.append([key,label,*[data['periods'][y].get(key) for y in YEARS],help_])
 meta=wb.create_sheet('Thong_tin');meta.append(['key','value'])
 for k,v in data['meta'].items():meta.append([k,safe_text(v)])
 if result is not None:
  summary=wb.create_sheet('Ket_qua');summary.append(['Chỉ số','Giá trị / ghi chú'])
  summary.append(['F-Score',result['f_score']]);summary.append(['M-Score',result['m_score']]);summary.append(['Bản tính','Ảnh chụp kết quả; nhập lại vào ứng dụng để tính khi thay số.'])
  for v in result['errors']+result['warnings']:summary.append(['Lưu ý',safe_text(v)])
  for title,key in [('F_Score','f_rows'),('M_Score','m_rows')]:
   sh=wb.create_sheet(title);rows=result[key]
   if rows:
    sh.append(list(rows[0]));
    for row in rows:sh.append([safe_text(v) for v in row.values()])
 for sh in wb:format_sheet(sh)
 out=BytesIO();wb.save(out);return out.getvalue()

def import_workbook(content):
 if len(content)>5*1024*1024:raise ValueError('Tệp vượt 5 MB.')
 try:
  with ZipFile(BytesIO(content)) as z:
   if sum(i.file_size for i in z.infolist())>25*1024*1024:raise ValueError('Nội dung giải nén vượt 25 MB.')
  wb=load_workbook(BytesIO(content),read_only=True,data_only=False)
 except (BadZipFile, OSError, KeyError) as exc:raise ValueError('Không đọc được XLSX. Dùng mẫu tải từ ứng dụng.') from exc
 try:
  if not {'Du_lieu','Thong_tin'}.issubset(wb.sheetnames):raise ValueError('Cần hai sheet Du_lieu và Thong_tin.')
  ws=wb['Du_lieu'];ms=wb['Thong_tin']
  for sheet in (ws,ms):
   if sheet.max_row is None or sheet.max_column is None:sheet.calculate_dimension(force=True)
  if ws.max_row>100 or ws.max_column>10 or ms.max_row>40 or ms.max_column>5:raise ValueError('Cấu trúc tệp quá lớn; dùng mẫu gốc.')
  rows=list(ws.values)
  if list(rows[0][:5]) != ['field','Chỉ tiêu','t','t1','t2']:raise ValueError('Không đổi tên 5 cột đầu trong Du_lieu.')
  d=blank();seen=set()
  for row in rows[1:]:
   key=row[0]
   if not key:continue
   if key not in FIELDS or key in seen:raise ValueError(f'field lạ hoặc trùng: {key}')
   seen.add(key)
   for i,y in enumerate(YEARS,2):
    value=row[i] if i<len(row) else None
    if value in (None,''):value=None
    elif isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):raise ValueError(f'{key}/{y}: nhập số Excel, không nhập công thức hoặc số dạng chữ.')
    d['periods'][y][key]=value
  if seen != set(FIELDS):raise ValueError('Thiếu dòng chỉ tiêu; dùng mẫu đầy đủ.')
  if list(next(ms.values)[:2])!=['key','value']:raise ValueError('Thong_tin cần cột key,value.')
  seen_meta=set()
  for row in list(ms.values)[1:]:
   key,value=row[:2]
   if not key:continue
   if key not in d['meta'] or key in seen_meta:raise ValueError(f'Thông tin lạ hoặc trùng: {key}')
   seen_meta.add(key)
   if key in BOOL_KEYS:
    if value is None:value=d['meta'][key]
    if not isinstance(value,bool):raise ValueError(f'{key}: dùng TRUE hoặc FALSE.')
   elif key=='year':
    if isinstance(value,bool) or not isinstance(value,(int,float)) or int(value)!=value or not 1990<=value<=2100:raise ValueError('Năm báo cáo phải là số nguyên 1990–2100.')
    value=int(value)
   else:
    if value is not None and not isinstance(value,str):raise ValueError(f'{key}: cần văn bản.')
    value=value or ''
    if value.startswith('='):raise ValueError('Không dùng công thức trong Thong_tin.')
   d['meta'][key]=value
  choices={'scope':['Hợp nhất','Riêng lẻ'],'unit':['VND','Nghìn VND','Triệu VND','Tỷ VND'],'sector':['Doanh nghiệp phi tài chính','Ngân hàng / Bảo hiểm / Chứng khoán'],'issuance':['Có','Không','Chưa xác định']}
  for k,values in choices.items():
   if d['meta'][k] not in values:raise ValueError(f'{k}: giá trị không hợp lệ.')
  if len(d['meta']['ticker'])>20 or len(d['meta']['source'])>2000:raise ValueError('Mã hoặc nguồn quá dài.')
  return d
 finally:wb.close()

# Sector-specific financial analysis. Annual data only, no synthetic total score.
FIN_COMMON={
 'assets':('Tổng tài sản cuối năm','Tổng tài sản cùng phạm vi với LNST.'),
 'equity':('Vốn chủ sở hữu cuối năm','Tổng vốn chủ sở hữu, không chỉ vốn của cổ đông mẹ nếu dùng LNST toàn tập đoàn.'),
 'net_income':('LNST toàn doanh nghiệp','Không trộn LNST cổ đông mẹ với tổng vốn/tài sản hợp nhất.'),
 'pretax_income':('Lợi nhuận trước thuế','Báo cáo KQKD năm, có thể âm.'),
 'price':('Giá cổ phiếu (VND/cp)','Giá nhập tay, đơn vị VND/cp dù số liệu BCTC dùng triệu/tỷ. Ghi ngày giá và nguồn.'),
 'eps':('EPS cơ bản năm (VND/cp)','EPS cơ bản đã công bố hoặc đã điều chỉnh nhất quán; không dùng LNST/số cổ phiếu cuối kỳ để thay EPS.'),
 'bvps':('BVPS (VND/cp)','Giá trị sổ sách cổ phần phổ thông sau loại lợi ích không kiểm soát/cổ phần ưu đãi nếu có. Dùng nguồn đã đối chiếu.')}
BANK_FIELDS={**FIN_COMMON,
 'loans':('Cho vay khách hàng gộp cuối năm','Trước dự phòng; cùng phạm vi với nợ nhóm 2 và nợ xấu.'),
 'npl':('Nợ xấu nhóm 3–5','Cộng nhóm 3,4,5 từ thuyết minh phân loại nợ; cùng mẫu số cho vay khách hàng.'),
 'group2':('Nợ cần chú ý nhóm 2','Không gộp vào nợ xấu nhóm 3–5.'),
 'loan_reserves':('Dự phòng rủi ro cho vay khách hàng','Số dư dự phòng cho khoản cho vay khách hàng cuối năm, nhập dương. Không cộng dự phòng chứng khoán/cam kết không cùng phạm vi.'),
 'deposits':('Tiền gửi khách hàng cuối năm','Không gồm tiền gửi/vay tổ chức tín dụng.'),
 'demand_deposits':('Tiền gửi không kỳ hạn khách hàng','Cùng phạm vi tiền gửi khách hàng để tính CASA.'),
 'net_interest':('Thu nhập lãi thuần trong năm','Báo cáo KQKD: thu nhập lãi và các khoản tương tự trừ chi phí lãi.'),
 'earning_assets':('Tài sản sinh lãi cuối năm','Xác định theo thuyết minh/IR, cùng định nghĩa hai năm. Không tự dùng tổng tài sản thay thế.'),
 'operating_income':('Tổng thu nhập hoạt động thuần (TOI)','Thu nhập lãi thuần cộng các khoản thu nhập ngoài lãi thuần; trước chi phí hoạt động và dự phòng.'),
 'operating_cost':('Chi phí hoạt động trong năm','Nhập dương; không cộng chi phí lãi hoặc chi phí dự phòng.'),
 'credit_cost':('Chi phí dự phòng tín dụng cho vay ròng','Chi phí trong kỳ trừ hoàn nhập, có thể âm. Cần cùng phạm vi cho vay khách hàng; không dùng tổng chi phí dự phòng nếu gồm tài sản khác.'),
 'car_reported':('CAR công bố (%)','Nhập 12.5 nếu CAR = 12,5%. Lấy từ công bố an toàn vốn, không suy ra từ vốn chủ/tài sản.'),
 'nim_reported':('NIM công bố (%) — nếu có','Nhập 3.2 nếu NIM = 3,2%. Chỉ đối chiếu khi cùng kỳ và cùng phạm vi; có thể dùng bình quân tháng/ngày khác cách 2 mốc trong công cụ.')}
BROKER_FIELDS={**FIN_COMMON,
 'liabilities':('Tổng nợ phải trả cuối năm','Không đồng nhất với nợ vay chịu lãi; không tự suy ra giới hạn pháp lý.'),
 'borrowings':('Tổng vay chịu lãi cuối năm','Vay ngắn hạn + dài hạn cùng phạm vi; không cộng tiền của nhà đầu tư.'),
 'margin_loans':('Dư nợ cho vay ký quỹ cuối năm','Chỉ cho vay margin, không cộng ứng trước nếu thu nhập margin không gồm ứng trước.'),
 'margin_income':('Thu nhập lãi cho vay ký quỹ năm','Chỉ lãi margin khớp dư nợ; không dùng toàn bộ thu nhập tiền gửi/cho vay.'),
 'operating_revenue':('Doanh thu hoạt động năm','Theo KQKD công ty chứng khoán, cùng cách phân loại hai năm; không nhầm với lợi nhuận thuần hoạt động.'),
 'brokerage_revenue':('Doanh thu môi giới năm','Doanh thu môi giới, không phải lợi nhuận môi giới.'),
 'fvtpl_gains':('Lãi tài sản FVTPL trong năm','Lãi ghi nhận trên KQKD, có thể gồm lãi đánh giá lại. Nhập dương.'),
 'fvtpl_losses':('Lỗ tài sản FVTPL trong năm','Nhập trị tuyệt đối dương, gồm phần lỗ theo cùng phạm vi với lãi FVTPL.'),
 'fvtpl_assets':('Tài sản FVTPL cuối năm','Giá trị ghi sổ cuối kỳ, không đồng nghĩa chỉ cổ phiếu.'),
 'liquid_assets':('Tiền và tiền gửi ngắn hạn của công ty','Chỉ nguồn tiền thuộc công ty, không cộng tiền khách hàng. Không coi FVTPL là tiền.'),
 'current_liabilities':('Tổng nợ ngắn hạn cuối năm','Dùng chỉ báo bao phủ tiền đơn giản, không phải kiểm tra thanh khoản pháp lý.'),
 'capital_ratio_reported':('Tỷ lệ vốn khả dụng công bố (%)','Nhập 250 nếu tỷ lệ = 250%. Lấy Báo cáo tỷ lệ an toàn tài chính; không suy ra từ BCTC hợp nhất.')}
FIN_FIELDS={'bank':BANK_FIELDS,'broker':BROKER_FIELDS}
FIN_NAMES={'bank':'Ngân hàng','broker':'Công ty chứng khoán'}
SIGNED_FIN={'equity','net_income','pretax_income','net_interest','operating_income','credit_cost','eps','bvps','nim_reported','car_reported','capital_ratio_reported'}

def financial_blank(kind):
 return {'kind':kind,'meta':{'ticker':'','year':2025,'unit':'Triệu VND','scope':'Hợp nhất','source':'','capital_source':'','price_source':'','demo':False},
 'periods':{y:{k:None for k in FIN_FIELDS[kind]} for y in YEARS}}

def financial_demo(kind):
 d=financial_blank(kind);d['meta'].update(ticker='DEMO',source='Dữ liệu giả lập, không phải số liệu doanh nghiệp.',demo=True)
 for y in ['t','t1']:
  d['periods'][y].update(assets=1000.,equity=100.,net_income=20.,pretax_income=25.,price=30000.,eps=3000.,bvps=20000.)
 if kind=='bank':
  for y in ['t','t1']:d['periods'][y].update(loans=600.,npl=12.,group2=6.,loan_reserves=18.,deposits=700.,demand_deposits=140.,net_interest=32.,earning_assets=800.,operating_income=50.,operating_cost=20.,credit_cost=6.,car_reported=12.,nim_reported=4.)
 else:
  for y in ['t','t1']:d['periods'][y].update(liabilities=900.,borrowings=700.,margin_loans=150.,margin_income=15.,operating_revenue=100.,brokerage_revenue=30.,fvtpl_gains=40.,fvtpl_losses=10.,fvtpl_assets=200.,liquid_assets=100.,current_liabilities=500.,capital_ratio_reported=250.)
 for k in ['assets','equity']+(['loans','earning_assets'] if kind=='bank' else ['margin_loans']):d['periods']['t2'][k]=d['periods']['t1'][k]
 return d

def financial_calculate(d):
 kind=d['kind'];fields=FIN_FIELDS[kind];p=d['periods'];errors=[];notes=[]
 for y in YEARS:
  for k,v in p[y].items():
   if k not in fields:errors.append(f'{y}: chỉ tiêu không được hỗ trợ.');continue
   if v is None:continue
   if isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v):errors.append(f'{y}: {fields[k][0]} phải là số hữu hạn.');continue
   if k not in SIGNED_FIN and v<0:errors.append(f'{y}: {fields[k][0]} không được âm.')
   if k=='assets' and v<=0:errors.append(f'{y}: tổng tài sản phải > 0.')
  def compare(part,total,label):
   a,b=p[y].get(part),p[y].get(total)
   if isinstance(a,(int,float)) and isinstance(b,(int,float)) and a>b+max(abs(b),1)*1e-8:errors.append(f'{y}: {label}.')
  if isinstance(p[y].get('equity'),(int,float)) and p[y]['equity']<=0:notes.append(f'{y}: vốn chủ sở hữu không dương; không diễn giải ROE/đòn bẩy như trường hợp vốn chủ dương.')
  if kind=='bank':
   compare('npl','loans','Nợ xấu vượt cho vay khách hàng');compare('group2','loans','Nợ nhóm 2 vượt cho vay khách hàng');compare('demand_deposits','deposits','Tiền gửi không kỳ hạn vượt tiền gửi khách hàng')
   a,b,c=(p[y].get(k) for k in ['npl','group2','loans'])
   if all(isinstance(v,(int,float)) for v in [a,b,c]) and a+b>c:errors.append(f'{y}: nợ nhóm 2 cộng nhóm 3–5 vượt cho vay khách hàng.')
  else:
   compare('borrowings','liabilities','Nợ vay vượt tổng nợ phải trả');compare('margin_loans','assets','Cho vay ký quỹ vượt tổng tài sản');compare('fvtpl_assets','assets','Tài sản FVTPL vượt tổng tài sản');compare('liquid_assets','assets','Tiền và tiền gửi vượt tổng tài sản');compare('current_liabilities','liabilities','Nợ ngắn hạn vượt tổng nợ')
   a,b,c=(p[y].get(k) for k in ['assets','equity','liabilities'])
   if all(isinstance(v,(int,float)) for v in [a,b,c]) and abs(a-b-c)>max(abs(a),1)*.005:notes.append(f'{y}: tổng tài sản không khớp vốn chủ + nợ (chênh > 0,5%); kiểm tra phạm vi/đơn vị.')
 if d['meta'].get('demo'):notes.append('Dữ liệu minh họa — không đại diện cho cổ phiếu thực tế.')
 if not d['meta'].get('source','').strip():notes.append('Chưa ghi nguồn BCTC, trang và thuyết minh.')
 if not d['meta'].get('price_source','').strip():notes.append('Chưa có ngày giá/nguồn EPS/BVPS. P/E, P/B nếu tính được vẫn cần đối chiếu thời điểm và điều chỉnh cổ phiếu.')
 notes.append('Chỉ dùng dữ liệu cả năm; ROA/ROE và lợi suất dùng bình quân hai mốc đầu/cuối năm, có thể khác chỉ số công bố bình quân nhiều mốc. Không tự xếp hạng tốt/xấu hoặc gán điểm F/M-Score cho nhóm tài chính.')
 out={'errors':errors,'notes':notes,'rows':[]}
 if errors:return out
 def get(y,k):
  v=p[y].get(k)
  if v is None:raise ValueError(f'Thiếu {fields[k][0]} ({y})')
  return v
 def div(a,b,label):
  if not math.isfinite(b) or b<=0:raise ValueError(f'{label}: mẫu số phải dương và hữu hạn')
  val=a/b
  if not math.isfinite(val):raise ValueError(f'{label}: kết quả vượt phạm vi số')
  return val
 def avg(y,k):return (get(y,k)+get('t1' if y=='t' else 't2',k))/2
 def ratio(y,n,den):return div(get(y,n),get(y,den),fields[den][0])
 def add(name,fn,unit,formula,note):
  vals=[];reasons=[]
  for y in ['t','t1']:
   try:vals.append(fn(y));reasons.append('')
   except ValueError as exc:vals.append(None);reasons.append(str(exc))
  diff=vals[0]-vals[1] if all(v is not None for v in vals) else None
  out['rows'].append({'Chỉ tiêu':name,'t':vals[0],'t1':vals[1],'Thay đổi':diff,'Đơn vị':unit,'Công thức':formula,'Ý nghĩa / Giới hạn':note,'Thiếu / Không xác định':'; '.join(f'{y}: {r}' for y,r in zip(['t','t1'],reasons) if r)})
 add('ROA',lambda y:div(get(y,'net_income'),avg(y,'assets'),'Tài sản bình quân'),'%', 'LNST / tài sản bình quân','Tổng LNST khớp tổng tài sản. Tăng không tự chứng minh chất lượng lợi nhuận tốt.')
 add('ROE',lambda y:div(get(y,'net_income'),avg(y,'equity'),'Vốn chủ bình quân'),'%', 'LNST / vốn chủ bình quân','Vốn chủ âm/không dương: không diễn giải. Đòn bẩy cao cũng có thể làm ROE cao.')
 add('P/E cơ bản',lambda y:div(get(y,'price'),get(y,'eps'),'EPS'),'lần','Giá VND/cp / EPS cơ bản VND/cp','EPS <= 0: không tính. Giá và EPS cần đồng nhất điều chỉnh; không kết luận rẻ chỉ từ P/E thấp.')
 add('P/B',lambda y:div(get(y,'price'),get(y,'bvps'),'BVPS'),'lần','Giá VND/cp / BVPS VND/cp','BVPS của cổ phần phổ thông, cùng cách điều chỉnh với giá. Không thay bằng tổng vốn hợp nhất nếu có lợi ích không kiểm soát.')
 def growth(k):return div(get('t',k)-get('t1',k),get('t1',k),'Giá trị năm trước')
 # Growth is current/prior; prior growth requires t-2 flow, not requested.
 def growth_row(name,k):
  try:v=growth(k);reason=''
  except ValueError as exc:v=None;reason=str(exc)
  out['rows'].append({'Chỉ tiêu':name,'t':v,'t1':None,'Thay đổi':None,'Đơn vị':'%','Công thức':'(t − t−1)/t−1, chỉ khi gốc > 0','Ý nghĩa / Giới hạn':'Không tính tăng trưởng năm trước vì không thu thập đầy đủ dòng tiền/lợi nhuận t−2. Gốc âm/0 không chuyển thành tăng trưởng phần trăm.','Thiếu / Không xác định':reason})
 growth_row('Tăng trưởng LNST','net_income')
 if kind=='bank':
  add('NPL cho vay khách hàng',lambda y:ratio(y,'npl','loans'),'%', 'Nợ nhóm 3–5 / cho vay khách hàng gộp','Cần đọc nợ tái cơ cấu, ngoại bảng, nợ đã xử lý; không chỉ nhìn tỷ lệ cuối kỳ.')
  add('Nợ nhóm 2 / cho vay',lambda y:ratio(y,'group2','loans'),'%', 'Nợ nhóm 2 / cho vay khách hàng gộp','Theo dõi nợ có thể chuyển nhóm; không gộp nhóm 2 vào NPL.')
  add('Bao phủ nợ xấu (LLR)',lambda y:ratio(y,'loan_reserves','npl'),'%', 'Dự phòng cho vay cuối kỳ / nợ xấu','NPL = 0 không tính tỷ lệ. Cần khớp phạm vi dự phòng với cho vay; tỷ lệ không chứng minh đủ dự phòng cho mọi tài sản.')
  add('CASA khách hàng',lambda y:ratio(y,'demand_deposits','deposits'),'%', 'Tiền gửi KKH khách hàng / tiền gửi khách hàng','Định nghĩa chỉ tiền gửi khách hàng, có thể khác CASA do ngân hàng công bố nếu phạm vi khác.')
  add('Cho vay / tiền gửi khách hàng',lambda y:ratio(y,'loans','deposits'),'%', 'Cho vay khách hàng gộp / tiền gửi khách hàng','Chỉ báo cơ cấu nguồn vốn; KHÔNG phải LDR pháp lý và không kiểm tra tuân thủ.')
  add('NIM tính bình quân hai mốc',lambda y:div(get(y,'net_interest'),avg(y,'earning_assets'),'Tài sản sinh lãi bình quân'),'%', 'Thu nhập lãi thuần năm / TS sinh lãi bình quân','Chỉ là cách tính hai mốc; có thể khác NIM bình quân ngày/tháng/quý công bố.')
  add('NIM công bố',lambda y:get(y,'nim_reported')/100,'%', 'Nhập từ tài liệu IR / công bố ngân hàng','Không tự thay NIM tính toán; đối chiếu phạm vi/kỳ/bình quân trước khi so sánh.')
  add('CIR',lambda y:ratio(y,'operating_cost','operating_income'),'%', 'Chi phí hoạt động / tổng thu nhập hoạt động thuần','Không gồm chi phí lãi hoặc dự phòng. CIR thấp hơn cần xem có khoản thu nhập bất thường không.')
  add('Chi phí tín dụng (xấp xỉ)',lambda y:div(get(y,'credit_cost'),avg(y,'loans'),'Cho vay bình quân'),'%', 'Chi phí dự phòng cho vay ròng / cho vay gộp bình quân','Chi phí trong kỳ, không phải số dư dự phòng. Hoàn nhập có thể làm âm.')
  add('CAR công bố',lambda y:get(y,'car_reported')/100,'%', 'Nhập từ công bố an toàn vốn','Không tính CAR bằng vốn chủ/tài sản. Không kiểm tra ngưỡng pháp lý; ghi nguồn/phạm vi và chuẩn áp dụng.')
  growth_row('Tăng trưởng cho vay khách hàng','loans');growth_row('Tăng trưởng tiền gửi khách hàng','deposits')
 else:
  add('Nợ vay / vốn chủ',lambda y:ratio(y,'borrowings','equity'),'lần','Vay chịu lãi / vốn chủ cuối kỳ','Không đồng nhất với tổng nợ phải trả; không tự kết luận vi phạm giới hạn pháp lý.')
  add('Tổng tài sản / vốn chủ',lambda y:ratio(y,'assets','equity'),'lần','Tổng tài sản / vốn chủ cuối kỳ','Chỉ báo đòn bẩy; không dùng thay tỷ lệ vốn khả dụng.')
  add('Dư nợ margin / vốn chủ',lambda y:ratio(y,'margin_loans','equity'),'lần','Cho vay ký quỹ / vốn chủ cuối kỳ','Cần đọc tập trung mã/tài sản bảo đảm, dự phòng và khả năng thu hồi. Không tự kiểm tra tuân thủ hạn mức.')
  add('Lợi suất cho vay margin (xấp xỉ)',lambda y:div(get(y,'margin_income'),avg(y,'margin_loans'),'Margin bình quân'),'%', 'Lãi margin năm / dư nợ margin bình quân hai mốc','Không cộng lãi tiền gửi/ứng trước nếu mẫu số chỉ là margin; không phải lãi suất hợp đồng.')
  add('Doanh thu môi giới / DT hoạt động',lambda y:ratio(y,'brokerage_revenue','operating_revenue'),'%', 'Doanh thu môi giới / doanh thu hoạt động','Là tỷ trọng doanh thu, không phải tỷ trọng lợi nhuận môi giới.')
  add('Lãi margin / DT hoạt động',lambda y:ratio(y,'margin_income','operating_revenue'),'%', 'Thu nhập lãi margin / doanh thu hoạt động','Không coi toàn bộ doanh thu cho vay/tiền gửi là lãi margin.')
  add('FVTPL ròng / DT hoạt động',lambda y:div(get(y,'fvtpl_gains')-get(y,'fvtpl_losses'),get(y,'operating_revenue'),'Doanh thu hoạt động'),'%', '(Lãi FVTPL − lỗ FVTPL) / doanh thu hoạt động','Có thể gồm đánh giá lại chưa thực hiện; không phải toàn bộ lợi nhuận tự doanh và các tỷ trọng không bắt buộc cộng bằng 100%.')
  add('Tài sản FVTPL / vốn chủ',lambda y:ratio(y,'fvtpl_assets','equity'),'lần','Giá trị FVTPL / vốn chủ cuối kỳ','FVTPL có thể gồm cổ phiếu, trái phiếu, chứng chỉ tiền gửi; cần xem cơ cấu và mức tập trung.')
  add('Bao phủ nợ ngắn hạn bằng tiền',lambda y:ratio(y,'liquid_assets','current_liabilities'),'lần','Tiền và tiền gửi NH thuộc công ty / nợ ngắn hạn','Không gồm tiền khách hàng; là chỉ báo đơn giản, không phải kết luận an toàn thanh khoản pháp lý.')
  add('Biên lợi nhuận trước thuế',lambda y:ratio(y,'pretax_income','operating_revenue'),'%', 'LNTT / doanh thu hoạt động','Phân loại doanh thu/lãi FVTPL cần nhất quán, không so ngang máy móc giữa các công ty.')
  add('Tỷ lệ vốn khả dụng công bố',lambda y:get(y,'capital_ratio_reported')/100,'%', 'Nhập từ Báo cáo tỷ lệ an toàn tài chính','Không suy ra từ vốn chủ. Ghi ngày, phạm vi riêng lẻ và căn cứ tính của báo cáo; không gắn ngưỡng tự động.')
  growth_row('Tăng trưởng dư nợ margin','margin_loans');growth_row('Tăng trưởng doanh thu hoạt động','operating_revenue')
 return out


def financial_export(data,result=None):
 wb=Workbook();ws=wb.active;ws.title='Tai_chinh';ws.append(['field','Chỉ tiêu','t','t1','t2','Hướng dẫn'])
 for k,(label,help_) in FIN_FIELDS[data['kind']].items():ws.append([k,label,*[data['periods'][y].get(k) for y in YEARS],help_])
 ms=wb.create_sheet('Thong_tin_TC');ms.append(['key','value']);ms.append(['kind',data['kind']])
 for k,v in data['meta'].items():ms.append([k,safe_text(v)])
 if result:
  sh=wb.create_sheet('Chi_tieu');rows=result['rows']
  if rows:
   sh.append(list(rows[0]))
   for row in rows:sh.append([safe_text(v) for v in row.values()])
   for i,row in enumerate(rows,2):
    if row['Đơn vị']=='%':
     for j in [2,3,4]:sh.cell(i,j).number_format='0.00%'
  note=wb.create_sheet('Luu_y');note.append(['Nội dung']);note.append(['Kết quả là ảnh chụp; nhập lại vào ứng dụng để tính sau khi sửa số. Thay đổi tỷ lệ là chênh điểm phần trăm khi hiển thị %.'])
  for n in result['errors']+result['notes']:note.append([safe_text(n)])
 for sh in wb:format_sheet(sh)
 if result and result['rows']:
  sh=wb['Chi_tieu']
  for i,row in enumerate(result['rows'],2):
   if row['Đơn vị']=='%':
    for j in [2,3,4]:sh.cell(i,j).number_format='0.00%'
 out=BytesIO();wb.save(out);return out.getvalue()


def financial_import(content,kind):
 if len(content)>5*1024*1024:raise ValueError('Tệp vượt 5 MB.')
 with ZipFile(BytesIO(content)) as z:
  if sum(i.file_size for i in z.infolist())>25*1024*1024:raise ValueError('Tệp giải nén quá lớn.')
 wb=load_workbook(BytesIO(content),read_only=True,data_only=False)
 try:
  if not {'Tai_chinh','Thong_tin_TC'}.issubset(wb.sheetnames):raise ValueError('Dùng mẫu Excel tài chính đúng nhóm, không dùng mẫu F/M-Score.')
  for sh in [wb['Tai_chinh'],wb['Thong_tin_TC']]:
   if sh.max_row is None or sh.max_column is None:sh.calculate_dimension(force=True)
   if sh.max_row>100 or sh.max_column>10:raise ValueError('Tệp không đúng kích thước mẫu.')
  rows=list(wb['Tai_chinh'].values);meta=list(wb['Thong_tin_TC'].values)
  if list(rows[0][:5])!=['field','Chỉ tiêu','t','t1','t2'] or list(meta[0][:2])!=['key','value']:raise ValueError('Không đổi tên cột trong mẫu.')
  md={};
  for row in meta[1:]:
   k,v=row[:2]
   if not k:continue
   if k in md:raise ValueError('Trùng thông tin: '+str(k))
   md[k]=v
  if md.pop('kind',None)!=kind:raise ValueError('Sai nhóm: tải và dùng mẫu của '+FIN_NAMES[kind])
  data=financial_blank(kind)
  if set(md)!=set(data['meta']):raise ValueError('Thiếu hoặc thừa trường thông tin.')
  for k,v in md.items():
   if k=='year':
    if isinstance(v,bool) or not isinstance(v,(int,float)) or int(v)!=v or not 1990<=v<=2100:raise ValueError('Năm phải là số nguyên 1990–2100.')
    v=int(v)
   elif k=='demo':
    if not isinstance(v,bool):raise ValueError('demo dùng TRUE/FALSE.')
   else:
    if v is None:v=''
    if not isinstance(v,str) or len(v)>3000 or v.startswith('='):raise ValueError('Thông tin phải là văn bản, không dùng công thức.')
   data['meta'][k]=v
  if data['meta']['unit'] not in ['VND','Nghìn VND','Triệu VND','Tỷ VND'] or data['meta']['scope'] not in ['Hợp nhất','Riêng lẻ']:raise ValueError('Đơn vị/phạm vi không hợp lệ.')
  seen=set()
  for row in rows[1:]:
   k=row[0]
   if not k:continue
   if k not in FIN_FIELDS[kind] or k in seen:raise ValueError('Chỉ tiêu lạ/trùng: '+str(k))
   seen.add(k)
   for i,y in enumerate(YEARS,2):
    v=row[i] if i<len(row) else None
    if v in (None,''):v=None
    elif isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v):raise ValueError(f'{k}/{y}: dùng số, không dùng công thức hoặc số dạng chữ.')
    data['periods'][y][k]=v
  if seen!=set(FIN_FIELDS[kind]):raise ValueError('Thiếu dòng chỉ tiêu.')
  return data
 finally:wb.close()


# User-supplied bank rules, revised to expose missing data and model limitations.
import html
import json
DASHBOARD_MODEL='bank-custom-1.0-adjusted'

def bank_dashboard_model(data):
 base=financial_calculate(data)
 out={'model_version':DASHBOARD_MODEL,'errors':base['errors'],'health_rows':[],'risk_rows':[], 'health_score':None,'known_health':0,'risk_score':None,'risk_raw':None,'metrics':[]}
 if base['errors']:return out
 rows={r['Chỉ tiêu']:r for r in base['rows']};p=data['periods']
 def m(name,y='t'):
  row=rows[name];v=row[y]
  if v is None:raise ValueError(row['Thiếu / Không xác định'] or 'Chưa đủ dữ liệu: '+name)
  return v
 def get(k,y='t'):
  v=p[y][k]
  if v is None:raise ValueError('Thiếu '+BANK_FIELDS[k][0]+f' ({y})')
  return v
 def ratio(a,b,label):
  if b<=0:raise ValueError(label+': mẫu số không dương')
  return a/b
 def nim(y):
  # Only use published NIM if BOTH years are provided, never mix methodologies.
  if p['t']['nim_reported'] is not None and p['t1']['nim_reported'] is not None:return get('nim_reported',y)/100
  return m('NIM tính bình quân hai mốc',y)
 nim_basis='NIM công bố hai năm' if all(p[y]['nim_reported'] is not None for y in ['t','t1']) else 'NIM tính từ TS sinh lãi bình quân hai mốc'
 criteria=[
  ('LNST năm t dương',lambda:(get('net_income'),0),lambda a,b:a>b,'LNST > 0'),
  ('ROA năm t > 1%',lambda:(m('ROA'),.01),lambda a,b:a>b,'ROA = LNST / TS bình quân; ngưỡng nội bộ 1%'),
  ('ROA cải thiện',lambda:(m('ROA'),m('ROA','t1')),lambda a,b:a>b,'ROA t > t−1; cần tài sản cuối t−2'),
  ('ROE năm t > 12%',lambda:(m('ROE'),.12),lambda a,b:a>b,'ROE = LNST / vốn chủ bình quân; ngưỡng nội bộ 12%'),
  ('Vốn chủ / tài sản > 10%',lambda:(ratio(get('equity'),get('assets'),'Tài sản'),.10),lambda a,b:a>b,'Chỉ là tỷ lệ kế toán, không phải CAR hoặc ngưỡng pháp lý'),
  ('Nợ xấu / cho vay < 2%',lambda:(m('NPL cho vay khách hàng'),.02),lambda a,b:a<b,'NPL nhóm 3–5 / cho vay khách hàng gộp'),
  ('Bao phủ nợ xấu ≥ 100%',lambda:(m('Bao phủ nợ xấu (LLR)'),1.),lambda a,b:a>=b,'Dự phòng cho vay / nợ xấu; NPL = 0 thì chưa xác định tỷ lệ'),
  ('Cho vay / tiền gửi trong 70–85%',lambda:(m('Cho vay / tiền gửi khách hàng'),None),lambda a,b:.70<=a<=.85,'Khoảng quy ước từ mẫu người dùng; KHÔNG phải LDR pháp lý'),
  ('NIM ổn định hoặc tăng',lambda:(nim('t'),nim('t1')),lambda a,b:a>=b,nim_basis)]
 for name,fn,compare,rule in criteria:
  try:
   current,reference=fn();passed=bool(compare(current,reference))
   row={'Tiêu chí':name,'Đạt':passed,'Điểm':int(passed),'Giá trị':current,'Đối chiếu':reference,'Quy tắc / Lý do':rule}
  except ValueError as exc:row={'Tiêu chí':name,'Đạt':None,'Điểm':None,'Giá trị':None,'Đối chiếu':None,'Quy tắc / Lý do':str(exc)}
  out['health_rows'].append(row)
 out['known_health']=sum(r['Điểm'] is not None for r in out['health_rows'])
 if out['known_health']==9:out['health_score']=sum(r['Điểm'] for r in out['health_rows'])
 def npl_growth():return ratio(get('npl')-get('npl','t1'),get('npl','t1'),'Nợ xấu năm gốc')
 components=[
  ('NPL / cho vay × 30',lambda:m('NPL cho vay khách hàng')*30),
  ('Phần tăng nợ xấu dương × 15',lambda:max(0,npl_growth())*15),
  ('Phần cho vay/tiền gửi vượt 85% × 20',lambda:max(0,m('Cho vay / tiền gửi khách hàng')-.85)*20),
  ('Thiếu bao phủ so với 100% × 15',lambda:max(0,1-min(m('Bao phủ nợ xấu (LLR)'),1.5))*15),
  ('Phần CIR vượt 40% × 20',lambda:max(0,m('CIR')-.40)*20)]
 for name,fn in components:
  try:
   v=fn()
   if not math.isfinite(v):raise ValueError('Giá trị vượt phạm vi số')
   row={'Thành phần':name,'Đóng góp':v,'Lý do':''}
  except ValueError as exc:row={'Thành phần':name,'Đóng góp':None,'Lý do':str(exc)}
  out['risk_rows'].append(row)
 if all(r['Đóng góp'] is not None for r in out['risk_rows']):
  out['risk_raw']=sum(r['Đóng góp'] for r in out['risk_rows']);out['risk_score']=min(10.,max(0.,out['risk_raw']))
 for name in ['NPL cho vay khách hàng','Cho vay / tiền gửi khách hàng','Bao phủ nợ xấu (LLR)','CIR']:
  out['metrics'].append({'Biến số':name,'Giá trị':rows[name]['t'],'Đơn vị':'%'})
 try:leverage=ratio(get('assets'),get('equity'),'Vốn chủ')
 except ValueError:leverage=None
 out['metrics'].append({'Biến số':'Tài sản / vốn chủ (bổ trợ, không vào FR)','Giá trị':leverage,'Đơn vị':'lần'})
 return out


def bank_dashboard_export(data):
 base=financial_calculate(data);model=bank_dashboard_model(data)
 wb=load_workbook(BytesIO(financial_export(data,base)))
 for title,rows in [('Bang_9_tieu_chi',model['health_rows']),('FR_thanh_phan',model['risk_rows'])]:
  ws=wb.create_sheet(title)
  if rows:
   ws.append(list(rows[0]))
   for r in rows:ws.append([safe_text(v) for v in r.values()])
  format_sheet(ws)
 ws=wb.create_sheet('Quy_uoc_bang');ws.append(['Mục','Giá trị'])
 for pair in [('Phiên bản mô hình',DASHBOARD_MODEL),('Số tiêu chí đủ dữ liệu',model['known_health']),('Điểm 9 tiêu chí (không phải F-Score)',model['health_score']),('FR hiển thị (chặn 0–10)',model['risk_score']),('FR tổng trước chặn',model['risk_raw']),('Ngưỡng quy ước',5),('Giới hạn','Quy tắc do người dùng cung cấp, có điều chỉnh; chưa hiệu chuẩn, không phải xác suất hoặc kết luận an toàn ngân hàng.')]:ws.append(pair)
 format_sheet(ws);out=BytesIO();wb.save(out);return out.getvalue()


def bank_collection_key(d):
 return '|'.join([d['meta']['ticker'].strip().upper(),str(d['meta']['year']),d['meta']['scope']])


def restore_bank_collection(content):
 if len(content)>5*1024*1024:raise ValueError('Bộ lưu vượt 5 MB.')
 def reject(value):raise ValueError('Không nhận NaN/Infinity trong bộ lưu.')
 obj=json.loads(content,parse_constant=reject)
 if not isinstance(obj,dict) or obj.get('format')!='StockAI_bank_collection_v1':raise ValueError('Sai định dạng bộ lưu.')
 records=obj.get('records')
 if not isinstance(records,list) or len(records)>100:raise ValueError('Bộ lưu tối đa 100 bản phân tích.')
 restored={}
 for r in records:
  if not isinstance(r,dict) or not isinstance(r.get('data'),dict):raise ValueError('Bản phân tích không hợp lệ.')
  d=r['data'];expected=financial_blank('bank')
  if d.get('kind')!='bank' or set(d.get('meta',{}))!=set(expected['meta']) or set(d.get('periods',{}))!=set(YEARS):raise ValueError('Sai cấu trúc ngân hàng.')
  if any(set(d['periods'][y])!=set(BANK_FIELDS) for y in YEARS):raise ValueError('Thiếu/thừa trường số liệu.')
  clean=financial_import(financial_export(d),'bank')
  if financial_calculate(clean)['errors']:raise ValueError('Bộ lưu có dữ liệu không hợp lệ.')
  if not re.fullmatch(r'[A-Z0-9]{1,15}',clean['meta']['ticker']):raise ValueError('Mã lưu không hợp lệ.')
  k=bank_collection_key(clean)
  if k in restored:raise ValueError('Trùng mã/năm/phạm vi trong bộ lưu.')
  restored[k]={'data':clean,'saved_at':str(r.get('saved_at',''))[:64]}
 return restored


# Plain-language interpretation added without changing v1.2 formulas or layout.

def sector_readout(data):
 base=financial_calculate(data);kind=data['kind'];p=data['periods']
 out={'health':'Chưa đủ dữ liệu để đánh giá','basis':'','checks':[],'flags':[], 'tested':0,'total_checks':0,
 'manipulation':'Chưa thể kết luận có hay không thao túng BCTC từ bộ số liệu này.'}
 if base['errors']:
  out['basis']='Có lỗi dữ liệu: '+'; '.join(base['errors']);return out
 rows={r['Chỉ tiêu']:r for r in base['rows']}
 def metric(name,y='t'):return rows[name][y]
 def test(label,fn,explain):
  out['total_checks']+=1
  try:
   values,condition=fn()
   if any(v is None for v in values):raise ValueError('Thiếu số liệu liên quan')
   flag=condition(*values);out['tested']+=1
   row={'Điều cần đối chiếu':label,'Kết quả':'Cần kiểm tra' if flag else 'Chưa thấy trong điều kiện này','Giải thích':explain}
   if flag:out['flags'].append(label+': '+explain)
  except (ValueError,TypeError,ZeroDivisionError):row={'Điều cần đối chiếu':label,'Kết quả':'Chưa đủ dữ liệu','Giải thích':explain}
  out['checks'].append(row)
 if kind=='bank':
  bd=bank_dashboard_model(data);score=bd['health_score'];out['basis']='Bộ 9 tiêu chí tham khảo của bản 1.2; không phải xếp hạng ngành hoặc CAR.'
  if score is not None:
   out['health']='Khỏe theo bộ tiêu chí tham khảo' if score>=8 else 'Trung bình / tín hiệu đan xen' if score>=5 else 'Yếu theo bộ tiêu chí tham khảo'
   out['basis']+=f' Đạt {score}/9 tiêu chí. Ngưỡng diễn giải nội bộ: 8–9, 5–7, 0–4.'
  if p['t']['equity'] is not None and p['t']['equity']<=0:out['health']='Có dấu hiệu suy yếu: vốn chủ không dương'
  test('Nợ nhóm 2 tăng tỷ trọng',lambda:([metric('Nợ nhóm 2 / cho vay'),metric('Nợ nhóm 2 / cho vay','t1')],lambda a,b:a>b),'Đọc phân loại nợ, nợ tái cơ cấu và thay đổi phạm vi báo cáo; đây là dấu hiệu rủi ro tín dụng, không phải bằng chứng thao túng.')
  test('Bao phủ nợ xấu giảm',lambda:([metric('Bao phủ nợ xấu (LLR)'),metric('Bao phủ nợ xấu (LLR)','t1')],lambda a,b:a<b),'Đối chiếu tăng nợ xấu, trích lập, hoàn nhập và nợ đã xử lý; giảm có thể có nguyên nhân hợp lệ.')
  test('LNST tăng trong khi chi phí dự phòng giảm',lambda:([p['t']['net_income'],p['t1']['net_income'],p['t']['credit_cost'],p['t1']['credit_cost']],lambda a,b,c,d:a>b and c<d),'Cần đọc cơ sở trích lập/hoàn nhập. Không tự kết luận giảm dự phòng là làm đẹp báo cáo.')
  def npl_outpaces(a,b,c,d):
   if b<=0 or d<=0:raise ValueError('Năm gốc không dương')
   return (a/b-1)>(c/d-1)
  test('NPL tăng nhanh hơn cho vay',lambda:([p['t']['npl'],p['t1']['npl'],p['t']['loans'],p['t1']['loans']],npl_outpaces),'Chỉ so tốc độ khi hai số năm gốc dương; cần xem cả giá trị tuyệt đối và chất lượng tín dụng.')
 else:
  ni,eq=p['t']['net_income'],p['t']['equity']
  directions=[('ROA',1),('ROE',1),('Nợ vay / vốn chủ',-1),('Bao phủ nợ ngắn hạn bằng tiền',1)]
  trend=[]
  for name,direction in directions:
   a,b=metric(name),metric(name,'t1')
   trend.append(None if a is None or b is None else direction*(a-b)>=0)
  out['basis']='Diễn giải từ LNST, vốn chủ và 4 xu hướng: ROA/ROE không giảm, nợ vay/vốn chủ không tăng, bao phủ tiền không giảm. Không tạo điểm F/M hoặc thang điểm ngành mới.'
  if ni is not None and eq is not None:
   if ni<0 or eq<=0:out['health']='Có dấu hiệu suy yếu: thua lỗ hoặc vốn chủ không dương'
   elif all(v is not None for v in trend):
    if ni>0 and all(trend):out['health']='Sinh lời dương, các xu hướng đang theo dõi không xấu đi'
    elif sum(v is False for v in trend)>=3:out['health']='Có nhiều xu hướng suy yếu cần kiểm tra'
    else:out['health']='Tín hiệu đan xen, chưa đủ cơ sở gọi là khỏe'
  test('Tỷ trọng FVTPL ròng tăng',lambda:([metric('FVTPL ròng / DT hoạt động'),metric('FVTPL ròng / DT hoạt động','t1')],lambda a,b:a>b),'Đọc lãi đã thực hiện/chưa thực hiện, phương pháp định giá và thanh khoản danh mục. Tự doanh tăng lãi không đồng nghĩa thao túng.')
  test('Đòn bẩy vay tăng',lambda:([metric('Nợ vay / vốn chủ'),metric('Nợ vay / vốn chủ','t1')],lambda a,b:a>b),'Kiểm tra nguồn vốn, kỳ hạn và tài sản bảo đảm; đây là dấu hiệu đòn bẩy, không phải bằng chứng sai BCTC.')
  test('Margin / vốn chủ tăng',lambda:([metric('Dư nợ margin / vốn chủ'),metric('Dư nợ margin / vốn chủ','t1')],lambda a,b:a>b),'Đối chiếu cơ cấu khách hàng, mã tài sản bảo đảm, dự phòng và khả năng thu hồi. Không tự kết luận vi phạm giới hạn pháp lý.')
  test('Bao phủ tiền giảm',lambda:([metric('Bao phủ nợ ngắn hạn bằng tiền'),metric('Bao phủ nợ ngắn hạn bằng tiền','t1')],lambda a,b:a<b),'Kiểm tra tiền thuộc công ty, không cộng tiền khách hàng; cần xem nguồn thanh khoản ngoài chỉ số này.')
 if out['flags']:
  out['manipulation']=f'Có {len(out["flags"])} điều kiện cần đối chiếu thêm; chưa đủ để kết luận hoặc ước tính xác suất thao túng BCTC.'
 elif out['tested']<out['total_checks']:
  out['manipulation']='Chưa đủ dữ liệu để hoàn tất các kiểm tra; không thể loại trừ rủi ro BCTC.'
 else:out['manipulation']='Chưa thấy các điều kiện cảnh báo đang theo dõi. Điều này KHÔNG chứng minh không có thao túng BCTC.'
 return out


def show_sector_readout(data):
 report=sector_readout(data)
 st.subheader('Diễn giải nhanh từ số liệu đã nhập')
 st.info('Sức khỏe tài chính: '+report['health'])
 st.caption(report['basis'])
 st.warning('Kiểm tra BCTC: '+report['manipulation'])
 st.caption(f'Đã thực hiện {report["tested"]}/{report["total_checks"]} điều kiện kiểm tra. Đây là kiểm tra tín hiệu tài chính, không phải mô hình phát hiện gian lận. Cần đọc báo cáo kiểm toán và thuyết minh.')
 if report['checks']:
  with st.expander('Vì sao có cảnh báo / cần đọc gì trong BCTC?'):
   st.dataframe(pd.DataFrame(report['checks']),hide_index=True,width='stretch')
 return report


def financial_export_with_readout(data,result=None):
 content=financial_export(data,result);wb=load_workbook(BytesIO(content));r=sector_readout(data)
 ws=wb.create_sheet('Dien_giai');ws.append(['Mục','Kết quả'])
 ws.append(['Sức khỏe',r['health']]);ws.append(['Cơ sở',r['basis']]);ws.append(['Thao túng BCTC',r['manipulation']]);ws.append(['Phạm vi','Bộ quy tắc tham khảo, không phải kết luận gian lận hoặc khuyến nghị đầu tư.'])
 format_sheet(ws)
 cs=wb.create_sheet('Can_doi_chieu')
 if r['checks']:
  cs.append(list(r['checks'][0]))
  for row in r['checks']:cs.append(list(row.values()))
 format_sheet(cs);out=BytesIO();wb.save(out);return out.getvalue()


def bank_dashboard_export_with_readout(data):
 content=bank_dashboard_export(data);wb=load_workbook(BytesIO(content));r=sector_readout(data)
 ws=wb.create_sheet('Dien_giai');ws.append(['Mục','Kết quả']);ws.append(['Sức khỏe',r['health']]);ws.append(['Cơ sở',r['basis']]);ws.append(['Thao túng BCTC',r['manipulation']]);format_sheet(ws)
 cs=wb.create_sheet('Can_doi_chieu')
 if r['checks']:
  cs.append(list(r['checks'][0]))
  for row in r['checks']:cs.append(list(row.values()))
 format_sheet(cs);out=BytesIO();wb.save(out);return out.getvalue()


def template_bytes(kind):
 name='Mau_Ngan_hang.xlsx' if kind=='bank' else 'Mau_Chung_khoan.xlsx'
 path=Path(__file__).parent/name
 return path.read_bytes() if path.exists() else financial_export(financial_blank(kind))


def bank_dashboard_page():
 st.markdown('''<style>
 .bank-head{padding:18px 0 22px;border-bottom:1px solid #334155;margin-bottom:20px;text-align:center}
 .bank-head h1{font-size:30px!important;color:#60a5fa!important;margin:0}.bank-head p{color:#94a3b8}
 .bank-card{background:#0b1220;color:#e2e8f0;border:1px solid #334155;border-radius:16px;padding:22px;margin:12px 0}
 .bank-card h3{color:#e2e8f0!important;font-size:17px!important}.bank-value{font-size:52px;line-height:1.2;font-weight:800;margin:10px 0}
 .bank-note{font-size:12px;color:#cbd5e1}.bank-row{padding:8px 0;border-bottom:1px solid #334155;font-size:12px;display:flex;justify-content:space-between;gap:10px}
 </style><div class="bank-head"><h1>Đánh giá sức khỏe &amp; cảnh báo ngân hàng</h1><p>9 tiêu chí tham khảo · FR-Index tùy chỉnh · Số liệu BCTC có nguồn</p></div>''',unsafe_allow_html=True)
 storage='finance_bank';revkey=storage+'_revision'
 if storage not in st.session_state:st.session_state[storage]=financial_blank('bank')
 if revkey not in st.session_state:st.session_state[revkey]=0
 if 'bank_saved' not in st.session_state:st.session_state.bank_saved={}
 def load(d):st.session_state[storage]=deepcopy(d);st.session_state[revkey]+=1
 with st.expander('Tạo mới, minh họa hoặc nhập Excel',expanded=False):
  a,b,c=st.columns(3)
  if a.button('Làm mới',key='board_new'):load(financial_blank('bank'));st.rerun()
  if b.button('Nạp minh họa',key='board_demo'):load(financial_demo('bank'));st.rerun()
  c.download_button('Tải mẫu Excel',template_bytes('bank'),'Mau_ngan_hang.xlsx',key='board_template')
  up=st.file_uploader('Chọn Excel ngân hàng',type=['xlsx'],key='board_upload')
  if st.button('Áp dụng Excel',disabled=up is None,key='board_apply'):
   try:load(financial_import(up.getvalue(),'bank'));st.rerun()
   except Exception as exc:st.error('Không nhập được: '+str(exc))
 data=st.session_state[storage];m=data['meta'];rev=st.session_state[revkey]
 def key(k):return f'board_{rev}_{k}'
 a,b,c=st.columns([1,1,1])
 m['ticker']=a.text_input('Mã ngân hàng',value=m['ticker'],max_chars=15,key=key('ticker')).strip().upper()
 m['year']=b.number_input('Năm báo cáo t',min_value=1990,max_value=2100,value=m['year'],key=key('year'))
 units=['VND','Nghìn VND','Triệu VND','Tỷ VND'];m['unit']=c.selectbox('Đơn vị BCTC',units,index=units.index(m['unit']),key=key('unit'))
 scopes=['Hợp nhất','Riêng lẻ'];m['scope']=st.selectbox('Phạm vi',scopes,index=scopes.index(m['scope']),key=key('scope'))
 m['source']=st.text_input('Nguồn BCTC và trang/thuyết minh',value=m['source'],max_chars=3000,key=key('source'))
 with st.expander('Nguồn bổ sung và nhãn minh họa'):
  m['capital_source']=st.text_input('Nguồn/kỳ/phạm vi CAR',value=m['capital_source'],max_chars=3000,key=key('capital_source'))
  m['price_source']=st.text_input('Ngày giá và nguồn EPS/BVPS',value=m['price_source'],max_chars=3000,key=key('price_source'))
  m['demo']=st.checkbox('Đây là dữ liệu minh họa',value=m['demo'],key=key('demo'))
 if m['demo']:st.warning('DỮ LIỆU MINHỌA — không phải BCTC thật của mã đang hiển thị.')
 st.caption('Các ngưỡng 1%, 12%, 10%, 2%, 70–85%, 5 điểm là quy ước nội bộ từ mẫu của bạn, không phải chuẩn ngành/pháp lý. Chỉ nhập BCTC cả năm; để trống khi chưa biết.')
 left,right=st.columns([1.4,1],gap='large')
 primary=['assets','equity','loans','deposits','npl','loan_reserves','net_income','net_interest','operating_income','operating_cost']
 def field(col,y,k):
  old=data['periods'][y][k]
  data['periods'][y][k]=col.number_input(BANK_FIELDS[k][0],value=float(old) if old is not None else None,format='%.4f',help=BANK_FIELDS[k][1],key=key(y+'_'+k))
 with left:
  st.subheader('Nhập số liệu hai năm')
  c1,c2=st.columns(2);c1.markdown(f'**NĂM t — {m["year"]}**');c2.markdown(f'**NĂM t−1 — {m["year"]-1}**')
  for k in primary:field(c1,'t',k);field(c2,'t1',k)
  with st.expander('Dữ liệu NIM và số đầu kỳ — cần để đủ 9 tiêu chí',expanded=True):
   st.caption('NIM dùng tài sản sinh lãi bình quân, không dùng dư nợ cho vay. Nhập NIM công bố cả hai năm hoặc bổ sung tài sản sinh lãi ba mốc. ROA so sánh cần tổng tài sản cuối t−2.')
   for k in ['earning_assets','nim_reported']:
    x,y=st.columns(2);x.caption('Năm t');y.caption('Năm t−1');field(x,'t',k);field(y,'t1',k)
   for k in ['assets','equity','loans','earning_assets']:field(st,'t2',k)
  with st.expander('Chỉ tiêu khác: CASA, nợ nhóm 2, CAR, định giá'):
   others=[k for k in BANK_FIELDS if k not in primary+['earning_assets','nim_reported']]
   x,y=st.columns(2);x.caption('Năm t');y.caption('Năm t−1')
   for k in others:field(x,'t',k);field(y,'t1',k)
 model=bank_dashboard_model(data);base=financial_calculate(data)
 with right:
  risk=model['risk_score'];health=model['health_score']
  risk_color='#fbbf24' if risk is None else '#fb7185' if risk>5 else '#60a5fa'
  risk_value='—' if risk is None else f'{risk:.1f}'
  risk_text='Chưa đủ dữ liệu' if risk is None else 'Vượt mốc quy ước 5 điểm' if risk>5 else 'Không vượt mốc quy ước 5 điểm'
  st.markdown(f'<div class="bank-card"><h3>Chỉ số cảnh báo tùy chỉnh (FR-Index)</h3><div class="bank-value" style="color:{risk_color}">{risk_value}<span style="font-size:16px;color:#94a3b8"> / 10</span></div><b>{risk_text}</b><p class="bank-note">Không phải xác suất gian lận hoặc kết luận ngân hàng an toàn. Mô hình chưa được hiệu chuẩn trên dữ liệu thực.</p></div>',unsafe_allow_html=True)
  for r in model['metrics']:
   v=r['Giá trị'];text='Chưa xác định' if v is None else f'{v*100:.2f}%' if r['Đơn vị']=='%' else f'{v:.2f} lần'
   st.markdown(f'<div class="bank-row"><span>{html.escape(r["Biến số"])}</span><b>{text}</b></div>',unsafe_allow_html=True)
  if risk is not None:
   st.caption(f'Tổng trước chặn: {model["risk_raw"]:.3f}. Hiển thị chặn trong [0,10]; trên 10 bị bão hòa. Đòn bẩy chỉ hiển thị bổ trợ, không có trọng số trong FR.')
  st.markdown(f'<div class="bank-card"><h3>9 tiêu chí sức khỏe tham khảo</h3><div class="bank-value" style="color:#34d399">{health if health is not None else "—"}<span style="font-size:24px;color:#94a3b8"> / 9</span></div><p class="bank-note">Đủ dữ liệu {model["known_health"]}/9 tiêu chí. Đây không phải Piotroski F-Score và không dùng để khẳng định “rất khỏe/an toàn”.</p></div>',unsafe_allow_html=True)
  for r in model['health_rows']:
   text='Chưa đủ' if r['Đạt'] is None else 'Đạt' if r['Đạt'] else 'Không đạt'
   color='#94a3b8' if r['Đạt'] is None else '#34d399' if r['Đạt'] else '#fbbf24'
   st.markdown(f'<div class="bank-row"><span>{html.escape(r["Tiêu chí"])}</span><b style="color:{color}">{text}</b></div>',unsafe_allow_html=True)
  for message in model['errors']:st.error(message)
  with st.expander('Chi tiết tiêu chí và từng đóng góp FR'):
   st.dataframe(pd.DataFrame(model['health_rows']),hide_index=True,width='stretch');st.dataframe(pd.DataFrame(model['risk_rows']),hide_index=True,width='stretch')
   st.caption('Sửa so với HTML: phần thiếu bao phủ = max(0, 1 − min(LLR,1.5)) × 15, không để dự phòng cao tạo đóng góp âm bù cảnh báo khác. Nếu một thành phần chưa xác định, không trả tổng FR.')
 show_sector_readout(data)
 with st.expander('Toàn bộ chỉ tiêu ngân hàng và lưu ý dữ liệu'):
  for n in base['notes']:st.warning(n)
  st.dataframe(pd.DataFrame(base['rows']),hide_index=True,width='stretch')
 a,b=st.columns(2)
 if a.button('Lưu / Cập nhật vào danh sách phiên',type='primary',key='board_save'):
  if not re.fullmatch(r'[A-Z0-9]{1,15}',m['ticker']):st.error('Nhập mã hợp lệ trước khi lưu.')
  elif base['errors']:st.error('Sửa lỗi số liệu trước khi lưu.')
  elif len(st.session_state.bank_saved)>=100 and bank_collection_key(data) not in st.session_state.bank_saved:st.error('Danh sách tối đa 100 bản.')
  else:
   st.session_state.bank_saved[bank_collection_key(data)]={'data':deepcopy(data),'saved_at':datetime.now().isoformat(timespec='seconds')}
   st.success('Đã lưu vào phiên. Hãy tải bộ lưu JSON hoặc Excel để giữ sau khi phiên kết thúc.')
 name=re.sub(r'[^A-Z0-9_-]','',m['ticker']) or 'CHUA_DAT_MA'
 b.download_button('Tải Excel gồm số liệu + điểm tham khảo',bank_dashboard_export_with_readout(data),f'{name}_{m["year"]}_bang_ngan_hang.xlsx',key='board_export')
 st.subheader('Danh sách ngân hàng đã lưu trong phiên')
 st.caption('Không tự lưu lâu dài trong trình duyệt và không phải kho dùng chung. Lưu theo mã + năm + phạm vi; cùng mã khác năm không ghi đè nhau. Tải bộ lưu để khôi phục lần sau.')
 book=st.session_state.bank_saved
 if book:
  table=[]
  for k,r in book.items():
   d=r['data'];s=bank_dashboard_model(d);table.append({'Mã':d['meta']['ticker'],'Năm':d['meta']['year'],'Phạm vi':d['meta']['scope'],'9 tiêu chí':s['health_score'],'FR tùy chỉnh':s['risk_score'],'Nhãn dữ liệu':'Minh họa' if d['meta']['demo'] else 'Người dùng nhập'})
  st.dataframe(pd.DataFrame(table),hide_index=True,width='stretch')
  selected=st.selectbox('Chọn bản đã lưu',list(book),key='board_saved_choice')
  a,b=st.columns(2)
  if a.button('Xem / Sửa',key='board_load'):load(book[selected]['data']);st.rerun()
  if b.button('Xóa khỏi danh sách phiên',key='board_delete'):del book[selected];st.rerun()
  backup={'format':'StockAI_bank_collection_v1','records':list(book.values())}
  st.download_button('Tải bộ lưu tất cả ngân hàng (.json)',json.dumps(backup,ensure_ascii=False,allow_nan=False,indent=2).encode('utf-8'),'Bo_luu_ngan_hang.json',key='board_backup')
 else:st.info('Chưa có bản lưu. Bạn có thể lưu cả bản chưa đủ điểm để nhập tiếp, nhưng điểm vẫn ghi chưa đủ dữ liệu.')
 with st.expander('Khôi phục bộ lưu JSON'):
  restored_file=st.file_uploader('Bộ lưu đã tải từ app (tối đa 5 MB)',type=['json'],key='board_restore_file')
  confirm=st.checkbox('Thay toàn bộ danh sách trong phiên bằng bộ lưu được chọn',key='board_restore_confirm')
  if st.button('Khôi phục danh sách',disabled=restored_file is None or not confirm,key='board_restore'):
   try:st.session_state.bank_saved=restore_bank_collection(restored_file.getvalue());st.rerun()
   except Exception as exc:st.error('Không khôi phục được: '+str(exc))


def financial_page(kind):
 if kind=='bank':
  bank_dashboard_page();return
 storage='finance_'+kind
 if storage not in st.session_state:st.session_state[storage]=financial_blank(kind)
 revkey=storage+'_revision'
 if revkey not in st.session_state:st.session_state[revkey]=0
 def load(d):st.session_state[storage]=deepcopy(d);st.session_state[revkey]+=1
 st.title('Phân tích '+FIN_NAMES[kind].lower())
 st.caption('Chỉ tiêu đặc thù ngành • Dữ liệu năm nhập tay/Excel • Không áp dụng F-Score/M-Score cho nhóm này')
 with st.expander('Tạo mới / minh họa / nhập Excel',expanded=True):
  a,b,c=st.columns(3)
  if a.button('Tạo bản mới',key=storage+'_new',width='stretch'):load(financial_blank(kind));st.rerun()
  if b.button('Nạp minh họa',key=storage+'_demo',width='stretch'):load(financial_demo(kind));st.rerun()
  c.download_button('Tải mẫu Excel',template_bytes(kind),f'Mau_{kind}.xlsx',key=storage+'_template',width='stretch')
  up=st.file_uploader('Chọn mẫu Excel đúng nhóm (tối đa 5 MB)',type=['xlsx'],key=storage+'_upload')
  if st.button('Áp dụng Excel',disabled=up is None,key=storage+'_apply'):
   try:load(financial_import(up.getvalue(),kind));st.rerun()
   except Exception as exc:st.error('Không nhập được: '+str(exc))
 data=st.session_state[storage];m=data['meta'];rev=st.session_state[revkey]
 def key(k):return f'{storage}_{rev}_{k}'
 a,b,c=st.columns(3)
 m['ticker']=a.text_input('Mã cổ phiếu',value=m['ticker'],max_chars=20,key=key('ticker')).strip().upper()
 m['year']=b.number_input('Năm báo cáo t',1990,2100,m['year'],key=key('year'))
 units=['VND','Nghìn VND','Triệu VND','Tỷ VND'];m['unit']=c.selectbox('Đơn vị số liệu BCTC',units,index=units.index(m['unit']),key=key('unit'))
 scopes=['Hợp nhất','Riêng lẻ'];m['scope']=st.selectbox('Phạm vi số liệu BCTC',scopes,index=scopes.index(m['scope']),key=key('scope'))
 m['source']=st.text_area('Nguồn BCTC, link và trang/thuyết minh',m['source'],max_chars=3000,key=key('source'))
 m['capital_source']=st.text_input('Nguồn, kỳ, phạm vi riêng lẻ/hợp nhất và chuẩn của CAR / tỷ lệ vốn khả dụng công bố',m['capital_source'],max_chars=3000,key=key('capital_source'))
 m['price_source']=st.text_input('Ngày giá t, ngày giá t−1; nguồn giá/EPS/BVPS và điều chỉnh cổ phiếu',m['price_source'],max_chars=3000,key=key('price_source'))
 m['demo']=st.checkbox('Đây là dữ liệu minh họa',m['demo'],key=key('demo'))
 st.info('Chỉ nhập BCTC cả năm. Ô trống là chưa biết; 0 là đã xác minh bằng 0. Giá/EPS/BVPS luôn VND/cp; CAR/NIM/vốn khả dụng công bố nhập theo %, ví dụ 12.5 = 12,5%. Đổi nhãn đơn vị không tự chuyển số.')
 tabs=st.tabs([f'Năm {m["year"]}',f'Năm {m["year"]-1}',f'Số đầu kỳ từ cuối {m["year"]-2}'])
 for i,y in enumerate(YEARS):
  with tabs[i]:
   keys=list(FIN_FIELDS[kind]) if y!='t2' else ['assets','equity']+(['loans','earning_assets'] if kind=='bank' else ['margin_loans'])
   cols=st.columns(2)
   for n,k in enumerate(keys):
    old=data['periods'][y][k]
    data['periods'][y][k]=cols[n%2].number_input(FIN_FIELDS[kind][k][0],value=float(old) if old is not None else None,format='%.4f',help=FIN_FIELDS[kind][k][1],key=key(y+'_'+k))
 result=financial_calculate(data)
 show_sector_readout(data)
 st.subheader('Bảng chỉ tiêu và so sánh hai năm')
 for n in result['errors']:st.error(n)
 for n in result['notes']:st.warning(n)
 if not m['capital_source'].strip():st.caption('Chưa ghi nguồn/phạm vi cho CAR hoặc vốn khả dụng; giá trị nhập tay chưa được xác minh.')
 if result['rows']:
  rows=[]
  def fmt(v,unit):return 'Chưa tính được' if v is None else f'{v*100:.2f}%' if unit=='%' else f'{v:.3f}'
  for r in result['rows']:
   row=dict(r);unit=r['Đơn vị'];row['t']=fmt(r['t'],unit);row['t1']=fmt(r['t1'],unit)
   row['Thay đổi']='—' if r['Thay đổi'] is None else (f'{r["Thay đổi"]*100:+.2f} điểm %' if unit=='%' else f'{r["Thay đổi"]:+.3f} lần');rows.append(row)
  st.dataframe(pd.DataFrame(rows),hide_index=True,width='stretch')
  st.caption('Tăng/giảm chỉ mô tả số liệu. Cần đọc nguyên nhân, cơ cấu tài sản và chất lượng lợi nhuận; diễn giải tham khảo không phải kết luận tuân thủ pháp luật hoặc khuyến nghị mua/bán.')
 name=re.sub(r'[^A-Z0-9_-]','',m['ticker']) or 'CHUA_DAT_MA'
 st.download_button('Tải dữ liệu và phân tích Excel',financial_export_with_readout(data,result),f'{name}_{m["year"]}_{kind}.xlsx',type='primary',key=storage+'_export')
 st.caption('Excel giữ số liệu và kết quả tại thời điểm tải; nhập lại tệp đúng nhóm để dùng tiếp. Bản độc lập không tự lưu dữ liệu lâu dài trên máy chủ.')

def replace(data):
 st.session_state.data=deepcopy(data)
 st.session_state.revision+=1

USER_GUIDE='## StockAI Độc lập 1.2 — bảng ngân hàng theo mẫu của bạn\n\n### Cập nhật web đã có\n\n1. Thay app.py và requirements.txt trong repository GitHub đang nối Streamlit, ở đúng đường dẫn app.py cũ. Bản này vẫn chỉ cần một app.py.\n2. Để có nền tối giống mẫu, thêm/thay tệp **.streamlit/config.toml** theo gói. Trên GitHub có thể chọn Add file → Create new file và gõ chính xác `.streamlit/config.toml`, dán nội dung tệp mẫu rồi Commit.\n3. Chờ app cập nhật hoặc Reboot. Chọn **Ngân hàng**: nhập hai năm ở trái, xem FR-Index và 9 tiêu chí ở phải; danh sách lưu ở bên dưới.\n4. Dùng mục Tạo mới, minh họa hoặc nhập Excel → Nạp minh họa để thử. Dữ liệu mặc định trống; minh họa không phải số liệu thật của TCB hay ngân hàng nào.\n\nCác phần F/M-Score cho doanh nghiệp phi tài chính, Công ty chứng khoán và Vĩ mô vẫn có. Bộ quy tắc ngân hàng không áp dụng cho công ty chứng khoán. Không cần Google Sheets, Apps Script hoặc Secrets trong bản độc lập.\n\n### Những điểm giữ theo mẫu HTML\n\n- Tông tối xanh, thẻ kết quả bên phải và hai cột năm t/t−1 bên trái.\n- Chín tiêu chí tham khảo, từng tiêu chí Đạt/Không đạt/Chưa đủ dữ liệu.\n- FR-Index tùy chỉnh với từng thành phần đóng góp, tổng trước chặn và số hiển thị 0–10.\n- Nhập mã, lưu/cập nhật, danh sách nhiều mã và xem/sửa/xóa bản lưu.\n\n### Những điểm đã điều chỉnh để tránh kết quả sai\n\n**1. Bộ quy tắc không phải mô hình nghiên cứu đã xác thực.** Các ngưỡng được lấy từ mẫu bạn gửi. App không gọi đó là Piotroski F-Score, xác suất gian lận, kết luận an toàn ngân hàng hay khuyến nghị mua/bán. Nhãn “không vượt 5” chỉ nói về mốc quy ước, không đồng nghĩa rủi ro thực thấp.\n\n**2. Thiếu dữ liệu không được coi là 0.** Nếu chưa đủ chín tiêu chí, tổng sức khỏe hiển thị “—”, kèm số tiêu chí đủ dữ liệu. Nếu một thành phần FR chưa xác định thì không trả tổng FR. Lưu được bản nháp chưa đủ điểm để làm tiếp.\n\n**3. ROA và ROE dùng bình quân hai mốc.** ROA t cần tài sản t và t−1; ROA t−1 cần t−1 và t−2. Tương tự vốn chủ bình quân cho ROE.\n\n**4. NIM đã sửa mẫu số.** Không lấy thu nhập lãi thuần/cho vay để gọi là NIM. Chọn một trong hai cách tự động, không trộn:\n- Nếu đã nhập NIM công bố của CẢ HAI năm: so sánh hai NIM công bố. Bạn cần đối chiếu cùng định nghĩa, kỳ và phạm vi.\n- Nếu không có đủ cả hai NIM công bố: dùng thu nhập lãi thuần/tài sản sinh lãi bình quân hai mốc cho cả hai năm. Cần tài sản sinh lãi t, t−1, t−2. Bình quân hai mốc có thể khác bình quân nhiều mốc công bố.\n\n**5. Cho vay/tiền gửi chỉ là tỷ lệ kế toán tham khảo.** Khoảng 70–85% trong tiêu chí là quy ước của mẫu bạn gửi, không phải kết luận LDR pháp lý. Tương tự vốn chủ/tài sản >10% không phải CAR hoặc tiêu chuẩn vốn pháp lý.\n\n**6. FR không có thành phần âm bù trừ cảnh báo khác.** Công thức hiện tại:\n\nFR thô = 30 × NPL/cho vay\n+ 15 × max(0, tốc độ tăng NPL)\n+ 20 × max(0, cho vay/tiền gửi − 0,85)\n+ 15 × max(0, 1 − min(bao phủ nợ xấu, 1,5))\n+ 20 × max(0, CIR − 0,40).\n\nFR hiển thị = min(10, max(0, FR thô)). FR > 5 là “vượt mốc quy ước”. Thẻ cũng hiển thị FR thô vì chặn trên 10 làm mất khác biệt giữa các trường hợp cao. Bao phủ >100% không tạo điểm âm giảm cảnh báo khác. Đòn bẩy được hiển thị bổ trợ nhưng KHÔNG có trọng số trong FR, đúng phạm vi công thức mẫu (không tự bổ sung trọng số).\n\nNPL năm trước bằng 0: không có tốc độ tăng phần trăm xác định, nên FR chưa tính được. NPL hiện tại bằng 0: tỷ lệ NPL là 0 nhưng bao phủ nợ xấu chưa xác định, không tự thay bằng 0 hoặc vô cực. Đây là giới hạn của bộ quy tắc này.\n\n### Chín tiêu chí đang dùng\n\n1. LNST t > 0.\n2. ROA t > 1%.\n3. ROA t > ROA t−1.\n4. ROE t > 12%.\n5. Vốn chủ cuối kỳ/tổng tài sản cuối kỳ > 10%.\n6. NPL nhóm 3–5/cho vay khách hàng gộp < 2%.\n7. Dự phòng cho vay khách hàng/NPL ≥ 100%.\n8. 70% ≤ cho vay khách hàng/tiền gửi khách hàng ≤ 85%.\n9. NIM t ≥ NIM t−1, cùng cách tính.\n\nTất cả là tiêu chí tham khảo nội bộ, chưa được hiệu chuẩn cho thị trường Việt Nam. Số tiêu chí đạt nhiều không thay thế đọc nợ tái cơ cấu, ngoại bảng, tập trung tín dụng, tài sản bảo đảm, nguồn vốn hoặc chất lượng lợi nhuận.\n\n### Lưu và khôi phục\n\nTrong Streamlit, bản này lưu danh sách ở **session_state**, không tự lưu lâu dài qua localStorage như HTML gốc.\n\n- Lưu/Cập nhật tạo bản trong phiên, khóa theo **mã + năm + hợp nhất/riêng lẻ**, nên không mất năm cũ khi thêm năm mới.\n- Đóng phiên/máy chủ khởi động lại có thể mất danh sách chưa tải. Để giữ lâu dài, bấm **Tải bộ lưu tất cả ngân hàng (.json)**.\n- Lần sau mở Khôi phục bộ lưu JSON, chọn tệp, xác nhận thay danh sách phiên rồi khôi phục. Số liệu được kiểm tra và điểm được tính lại, không tin vào điểm cũ trong tệp.\n- Danh sách tối đa 100 bản, tệp JSON tối đa 5 MB. Bộ khôi phục thay TOÀN BỘ danh sách phiên, không tự trộn.\n- Tải Excel từng bài để lưu số liệu, toàn bộ chỉ tiêu ngành, bảng 9 tiêu chí, thành phần FR và ghi chú quy ước. Tệp Excel đó nhập lại bằng mục Excel ngân hàng được.\n- Xóa chỉ xóa bản trong phiên hiện tại, không xóa các tệp bạn đã tải.\n- Đây không phải kho cộng đồng dùng chung; nếu muốn lưu bài của các thành viên trên máy chủ, cần tích hợp vào bản cộng đồng có cơ sở dữ liệu.\n\n### Cách nhập số đúng\n\nChỉ nhập BCTC cả năm, cùng đơn vị tiền, cùng phạm vi hợp nhất/riêng lẻ. Chi phí/dự phòng số dư nhập dương; LNST và thu nhập lãi thuần có thể âm. NPL là nhóm 3–5; cho vay khách hàng gộp trước dự phòng; tiền gửi là tiền gửi khách hàng, không gộp tiền gửi tổ chức tín dụng. Dự phòng cho vay cần cùng phạm vi với NPL.\n\nPhần mở rộng vẫn có CASA, nợ nhóm 2, chi phí tín dụng, CAR công bố và P/E/P/B. Giá/EPS/BVPS dùng VND/cp bất kể đơn vị BCTC. CAR/NIM công bố nhập 12.5 nếu là 12,5%. Ghi ngày và nguồn; không tự suy ra CAR từ tỷ lệ vốn chủ/tài sản.\n\nExcel dùng hai sheet Tai_chinh và Thong_tin_TC. Không đổi field/cột, không nhập công thức hay số dạng chữ. Các sheet điểm là ảnh chụp kết quả tại thời điểm tải; sửa số xong nhập lại app để tính. Mẫu ngân hàng và chứng khoán không thay thế nhau.\n\n### Giới hạn kiểm chứng\n\nỨng dụng chưa đọc PDF/cào BCTC, chưa kiểm tra nguồn nhập tay và chưa hiệu chuẩn FR/9 tiêu chí trên dữ liệu thực. Các số liệu minh họa giả lập chỉ để kiểm tra công thức. Không áp dụng điểm ngân hàng cho công ty chứng khoán.\n\n\n## Bổ sung theo yêu cầu: giữ nền bản 1.2, thêm mẫu và diễn giải\n\nĐây vẫn là bản độc lập StockAI_DocLap_v1_2, giữ bố cục, các phân hệ, công thức và bộ lưu JSON cá nhân. Không tích hợp R2, Supabase hoặc đăng nhập Google trong gói này. Chỉ thêm mẫu Excel và diễn giải sức khỏe/các điều cần kiểm tra.\n\n### Cập nhật web\n\nThay app.py và requirements.txt trong GitHub đang nối Streamlit. Đặt thêm **Mau_Ngan_hang.xlsx** và **Mau_Chung_khoan.xlsx** cùng thư mục với app.py. Giữ cấu hình giao diện tối .streamlit/config.toml của bản 1.2.\n\nNếu chỉ muốn dùng mẫu và chưa muốn cập nhật mã, hai tệp XLSX này cũng nhập được vào bản 1.2 cũ; phần diễn giải mới cần app.py trong gói này. Không đưa cả thư mục lồng vào sai đường dẫn app.py của web.\n\n### Người dùng chỉ cần làm gì?\n\n1. Chọn Ngân hàng hoặc Công ty chứng khoán, bấm Tải mẫu Excel.\n2. Sheet Tai_chinh: chỉ điền số ở các ô t, t1, t2 được tô màu; giữ nguyên field, tên cột và dòng.\n3. Sheet Thong_tin_TC: điền mã, năm, đơn vị, hợp nhất/riêng lẻ và nguồn BCTC. Các ô metadata nhập trên app cũng được sau khi nạp mẫu.\n4. Tải lên tệp → Áp dụng Excel. Ứng dụng tính tỷ lệ và hiển thị diễn giải.\n5. Tải kết quả Excel để lưu. Đổi số trong Excel kết quả không tự tính lại; phải nhập lại app.\n\nMẫu trống không chứa số liệu của ngân hàng/công ty thực tế. Dùng Nạp minh họa trong app nếu muốn xem thử có số. Ghi đúng đơn vị: các khoản BCTC cùng đơn vị; giá/EPS/BVPS luôn VND/cp. CAR/NIM/vốn khả dụng công bố nhập theo % (12.5 = 12,5%).\n\n### Ngân hàng: “khỏe/yếu” được hiểu thế nào?\n\nGiữ nguyên 9 tiêu chí/FR của bản 1.2. Thêm diễn giải nội bộ:\n- 8–9/9: Khỏe theo bộ tiêu chí tham khảo.\n- 5–7/9: Trung bình/tín hiệu đan xen.\n- 0–4/9: Yếu theo bộ tiêu chí tham khảo.\n- Chưa đủ tiêu chí: chưa đủ dữ liệu để đánh giá.\n- Vốn chủ không dương: cảnh báo suy yếu riêng, không dùng điểm tổng để che cảnh báo này.\n\nĐây là quy tắc do người dùng cung cấp, chưa hiệu chuẩn; không phải xếp hạng tín nhiệm, chuẩn pháp lý, F-Score hay dự báo lợi nhuận. FR vẫn là cảnh báo tài chính tùy chỉnh, không phải công cụ xác định gian lận.\n\n### Chứng khoán: không áp dụng điểm ngân hàng\n\nKhông tự áp bộ 9 tiêu chí ngân hàng sang CTCK. App diễn giải từ LNST/vốn chủ và bốn xu hướng: ROA, ROE không giảm; vay/vốn chủ không tăng; bao phủ tiền không giảm.\n- Lỗ hoặc vốn chủ không dương: có dấu hiệu suy yếu.\n- LNST dương, đủ bốn xu hướng và không xu hướng nào xấu đi: sinh lời dương, các xu hướng theo dõi không xấu đi.\n- Từ ba xu hướng xấu đi: có nhiều xu hướng suy yếu cần kiểm tra.\n- Tình huống còn lại: đan xen hoặc thiếu dữ liệu.\n\nCác xu hướng chỉ so với năm trước, không chứng minh mức tuyệt đối đã an toàn hay công ty “khỏe” theo chuẩn ngành. Không coi đòn bẩy/margin giảm luôn là quyết định kinh doanh tốt.\n\n### Có kết luận thao túng BCTC được không?\n\n**Không. Chỉ nhập các số hiện có chưa đủ để kết luận có hoặc không thao túng BCTC**, cũng không đủ để tạo xác suất gian lận đã kiểm chứng cho ngân hàng/CTCK.\n\nApp thêm bảng “Điều cần đối chiếu”:\n- Ngân hàng: nợ nhóm 2 tăng tỷ trọng; bao phủ NPL giảm; LNST tăng trong khi chi phí dự phòng giảm; NPL tăng nhanh hơn cho vay.\n- CTCK: tỷ trọng FVTPL ròng tăng; đòn bẩy vay tăng; margin/vốn chủ tăng; bao phủ tiền giảm.\n\nĐây chủ yếu là các dấu hiệu rủi ro tài chính hoặc cơ cấu lợi nhuận cần đọc thêm, KHÔNG phải bằng chứng làm sai báo cáo. Khi không xuất hiện cảnh báo, app vẫn ghi rõ không loại trừ thao túng.\n\nMuốn đánh giá sâu hơn phải đọc ý kiến kiểm toán, phân loại nợ/tài sản, căn cứ trích lập, lãi thực hiện/chưa thực hiện, giao dịch liên quan, điều chỉnh hồi tố và thuyết minh. App chưa tự đọc hay xác minh các tài liệu đó.\n\n### Lưu ý khi giữ bản cũ\n\nCác hướng dẫn phần nền của 1.2 vẫn có hiệu lực. Đoạn nền nói chưa gán nhãn khỏe/yếu được bổ sung bởi diễn giải nội bộ nêu trên; không có công thức số nào của 1.2 bị đổi. Không cần nâng sang bản R2 nếu bạn chưa muốn lưu cộng đồng.\n'

def main():
 st.set_page_config(page_title='StockAI | Phân tích BCTC',page_icon='📊',layout='wide')
 st.markdown('''<style>.block-container{max-width:1200px;padding-top:2rem}h1,h2,h3{color:#93c5fd}div[data-testid="stMetric"]{background:#0b1220;color:#e2e8f0;padding:16px;border-radius:12px}.stButton>button{border-radius:8px} </style>''',unsafe_allow_html=True)
 if 'data' not in st.session_state:st.session_state.data=blank()
 if 'revision' not in st.session_state:st.session_state.revision=0
 
 st.sidebar.title('NHÀ ĐẦU TƯ\nTHÔNG THÁI')
 st.sidebar.caption('Công cụ đọc hiểu BCTC • Phiên bản độc lập 1.2')
 page=st.sidebar.radio('Chức năng',['Phân tích BCTC','Ngân hàng','Công ty chứng khoán','Vĩ mô thủ công','Hướng dẫn'])
 st.sidebar.info('Dùng ngay, không cần tài khoản. Phiên làm việc không phải nơi lưu lâu dài; hãy tải Excel trước khi thoát.')
 if page=='Phân tích BCTC':
  st.title('Phân tích chất lượng tài chính')
  st.caption('Nhập dữ liệu có nguồn • Xem từng phép tính • Tải về để lưu và dùng tiếp')
  with st.expander('Bắt đầu: dữ liệu mới, minh họa hoặc nhập Excel',expanded=True):
   c1,c2,c3=st.columns(3)
   with c1:
    if st.button('Tạo bản mới',width='stretch'):replace(blank());st.rerun()
   with c2:
    if st.button('Nạp dữ liệu minh họa',width='stretch'):replace(demo());st.rerun()
   with c3:
    template=Path(__file__).parent/'Mau_nhap_BCTC.xlsx'
    if template.exists():st.download_button('Tải mẫu Excel trống',template.read_bytes(),'Mau_nhap_BCTC.xlsx',width='stretch')
    else:st.download_button('Tải mẫu Excel trống',export_workbook(blank()),'Mau_nhap_BCTC.xlsx',width='stretch')
   uploaded=st.file_uploader('Nhập mẫu Excel hoặc bản đã tải trước đó (.xlsx, tối đa 5 MB)',type=['xlsx'])
   if st.button('Áp dụng tệp Excel',disabled=uploaded is None):
    try:
     replace(import_workbook(uploaded.getvalue()));st.rerun()
    except Exception as exc:st.error(f'Không nhập được tệp: {exc}')
  d=st.session_state.data;m=d['meta'];rev=st.session_state.revision
  def key(name):return f'{rev}_{name}'
  st.subheader('1. Thông tin báo cáo')
  a,b,c=st.columns(3)
  m['ticker']=a.text_input('Mã cổ phiếu / doanh nghiệp',value=m['ticker'],max_chars=20,key=key('ticker')).strip().upper()
  m['year']=b.number_input('Năm báo cáo t',min_value=1990,max_value=2100,value=m['year'],key=key('year'))
  units=['VND','Nghìn VND','Triệu VND','Tỷ VND'];m['unit']=c.selectbox('Đơn vị toàn bộ số tiền',units,index=units.index(m['unit']),key=key('unit'))
  a,b=st.columns(2)
  scopes=['Hợp nhất','Riêng lẻ'];m['scope']=a.selectbox('Phạm vi báo cáo',scopes,index=scopes.index(m['scope']),key=key('scope'))
  sectors=['Doanh nghiệp phi tài chính','Ngân hàng / Bảo hiểm / Chứng khoán'];m['sector']=b.selectbox('Nhóm doanh nghiệp',sectors,index=sectors.index(m['sector']),key=key('sector'))
  m['source']=st.text_area('Nguồn BCTC, link và số trang/thuyết minh',value=m['source'],max_chars=2000,key=key('source'))
  m['demo']=st.checkbox('Bộ dữ liệu minh họa / chưa phải số liệu thực tế',value=m['demo'],key=key('demo'))
  if m['demo']:st.warning('DỮ LIỆU MINHỌA — không phải kết quả phân tích một cổ phiếu thực.')
  st.subheader('2. Số liệu hai năm và tài sản đầu kỳ')
  st.caption('Để trống nếu chưa biết. Số 0 có nghĩa bạn đã kiểm tra và khoản mục bằng 0. Nhập số trực tiếp, không gõ dấu phân cách hàng nghìn. Chi phí nhập dương; LNST và CFO có thể âm.')
  tabs=st.tabs([f'Năm t: {m["year"]}',f'Năm t−1: {m["year"]-1}',f'Tài sản t−2: {m["year"]-2}'])
  for index,y in enumerate(['t','t1','t2']):
   with tabs[index]:
    fields=list(FIELDS) if y!='t2' else ['assets']
    left,right=st.columns(2)
    for i,k in enumerate(fields):
     with (left if i%2==0 else right):
      old=d['periods'][y].get(k)
      d['periods'][y][k]=st.number_input(FIELDS[k][0],value=float(old) if old is not None else None,format='%.4f',help=FIELDS[k][1],key=key(f'{y}_{k}'))
  st.subheader('3. Kiểm tra thuyết minh')
  choices=['Chưa xác định','Không','Có'];m['issuance']=st.selectbox('Có phát hành thêm cổ phần phổ thông để huy động vốn trong năm t?',choices,index=choices.index(m['issuance']),key=key('issuance'),help='Không kết luận chỉ từ thay đổi số cổ phiếu. Chia tách/cổ tức cổ phiếu không đồng nghĩa phát hành huy động vốn.')
  m['f_income_proxy']=st.checkbox('F-Score: đang dùng LNST thay lợi nhuận trước khoản bất thường',value=m['f_income_proxy'],key=key('fp'))
  m['m_income_proxy']=st.checkbox('M-Score: đang dùng LNST thay lợi nhuận hoạt động tiếp tục',value=m['m_income_proxy'],key=key('mp'))
  if st.button('Sao chép LNST năm t sang lợi nhuận hoạt động tiếp tục (bản xấp xỉ)'):
   d['periods']['t']['continuing_income']=d['periods']['t']['income'];m['m_income_proxy']=True;replace(d);st.rerun()
  st.subheader('4. Kết quả cập nhật theo đầu vào')
  result=calculate(d)
  for msg in result['errors']:st.error(msg)
  for msg in result['warnings']:st.warning(msg)
  f=result['f_score'];ms=result['m_score']
  a,b=st.columns(2)
  a.metric('F-Score' + (' (điều chỉnh)' if m['f_income_proxy'] else ''),f'{f}/9' if f is not None else 'Chưa đủ dữ liệu')
  b.metric('M-Score 8 biến' + (' (xấp xỉ)' if m['m_income_proxy'] else ''),f'{ms:.3f}' if ms is not None else 'Chưa tính được')
  if f is not None:
   st.info('F-Score đếm số tín hiệu thuận lợi. Điểm cao hơn có nhiều tiêu chí đạt hơn; không thay thế định giá, triển vọng hay kiểm tra ngành.')
  if result['f_rows']:
   st.markdown('**Chín tiêu chí F-Score**');st.dataframe(pd.DataFrame(result['f_rows']),hide_index=True,width='stretch')
   known=[r['Điểm'] for r in result['f_rows'] if r['Điểm'] is not None]
   if f is None:st.caption(f'Đã tính {len(known)}/9 tiêu chí. Không cộng điểm một phần thành F-Score hoàn chỉnh.')
  if result['m_rows']:
   st.markdown('**Tám biến Beneish M-Score**');st.dataframe(pd.DataFrame(result['m_rows']),hide_index=True,width='stretch')
   st.caption('M = −4,84 + 0,920 DSRI + 0,528 GMI + 0,404 AQI + 0,892 SGI + 0,115 DEPI − 0,172 SGAI − 0,327 LVGI + 4,679 TATA.')
  if ms is not None:
   st.info('M-Score là chỉ báo sàng lọc dấu hiệu thao túng lợi nhuận. Bản này hiển thị số và phép tính, không tự gắn nhãn gian lận hay xác suất; chưa hiệu chuẩn ngưỡng cho dữ liệu Việt Nam.')
  st.subheader('5. Lưu để dùng tiếp')
  st.caption('Tệp Excel gồm dữ liệu nhập, nguồn, thông tin báo cáo và bảng kết quả tại thời điểm tải. Có thể nhập lại tệp để tiếp tục; Excel không tự tính lại khi bạn sửa số trong tệp.')
  safe_name=re.sub(r'[^A-Z0-9_-]','',m['ticker']) or 'CHUA_DAT_MA'
  st.download_button('Tải dữ liệu + kết quả Excel',export_workbook(d,result),f'{safe_name}_{m["year"]}_phan_tich.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
 elif page in ['Ngân hàng','Công ty chứng khoán']:
  financial_page('bank' if page=='Ngân hàng' else 'broker')
 elif page=='Vĩ mô thủ công':
  st.title('Bảng theo dõi vĩ mô thủ công')
  st.caption('Quy tắc tham khảo do người dùng nhập; không tự lấy dữ liệu hoặc dự báo thị trường.')
  source=st.text_area('Nguồn, ngày số liệu và ghi chú',key='macro_source')
  choices=['Chưa chọn','Thuận lợi','Trung tính','Bất lợi']
  labels=['Lãi suất và thanh khoản','Tỷ giá USD/VND','Sức mạnh USD (DXY)','Điều hành tiền tệ']
  values=[st.selectbox(x,choices,key=f'macro_{i}') for i,x in enumerate(labels)]
  if 'Chưa chọn' in values:st.info('Chọn đủ bốn yếu tố để xem tổng hợp.')
  else:
   status='Có yếu tố bất lợi' if 'Bất lợi' in values else 'Các yếu tố cùng thuận lợi' if all(v=='Thuận lợi' for v in values) else 'Các yếu tố chưa đồng thuận'
   st.info(status+'. Cần kiểm chứng dữ liệu, kết hợp định giá và mức chịu rủi ro; không quy đổi thành khuyến nghị all-in hoặc tỷ lệ tiền mặt cố định.')
   report='ĐÁNH GIÁ VĨ MÔ THỦ CÔNG\n'+ '\n'.join(f'{k}: {v}' for k,v in zip(labels,values))+'\nTổng hợp: '+status+'\nNguồn/ghi chú: '+source
   st.download_button('Tải ghi chú vĩ mô',report.encode('utf-8-sig'),'ghi_chu_vi_mo.txt')
 else:
  st.title('Hướng dẫn sử dụng')
  st.markdown(USER_GUIDE)
 st.divider();st.caption('StockAI Độc lập 1.2 • Công cụ hỗ trợ đọc BCTC. Dữ liệu do người dùng cung cấp, chưa được đối chiếu tự động với báo cáo gốc.')

if __name__=="__main__":
 main()
