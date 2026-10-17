"""Workbook interchange used by the application at runtime."""
from io import BytesIO
import math
from zipfile import ZipFile, BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from engine import FIELDS, YEARS, blank

BOOL_KEYS={'demo','f_income_proxy','m_income_proxy'}
def safe_text(v):
 return "'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v

def format_sheet(ws):
 ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='123047')
 for col in ws.columns:
  letter=col[0].column_letter;ws.column_dimensions[letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col)+2))
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(cell.value,(int,float)):cell.number_format='#,##0.000;[Red](#,##0.000);0'
  ws.row_dimensions[row[0].row].height=32

def export_workbook(data, result=None):
 wb=Workbook();ws=wb.active;ws.title='Du_lieu'
 ws.append(['field','Chỉ tiêu','t','t1','t2','Hướng dẫn'])
 for key,(label,help_) in FIELDS.items():ws.append([key,label,*[data['periods'][y].get(key) for y in YEARS],help_])
 meta=wb.create_sheet('Thong_tin');meta.append(['key','value'])
 for k,v in data['meta'].items():meta.append([k,safe_text(v)])
 if result is not None:
  summary=wb.create_sheet('Ket_qua');summary.append(['Chỉ số','Giá trị / ghi chú'])
  summary.append(['F-Score',result['f_score']]);summary.append(['M-Score',result['m_score']]);summary.append(['Bản tính','Ảnh chụp kết quả; nhập lại vào ứng dụng để tính khi thay số.'])
  for v in result['errors']+result['warnings']:summary.append(['Lưu ý',safe_text(v)])
  for title,key in [('F_Score','f_rows'),('M_Score','m_rows')]:
   sh=wb.create_sheet(title);rows=result[key]
   if rows:
    sh.append(list(rows[0]));
    for row in rows:sh.append([safe_text(v) for v in row.values()])
 for sh in wb:format_sheet(sh)
 out=BytesIO();wb.save(out);return out.getvalue()

def import_workbook(content):
 if len(content)>5*1024*1024:raise ValueError('Tệp vượt 5 MB.')
 try:
  with ZipFile(BytesIO(content)) as z:
   if sum(i.file_size for i in z.infolist())>25*1024*1024:raise ValueError('Nội dung giải nén vượt 25 MB.')
  wb=load_workbook(BytesIO(content),read_only=True,data_only=False)
 except (BadZipFile, OSError, KeyError) as exc:raise ValueError('Không đọc được XLSX. Dùng mẫu tải từ ứng dụng.') from exc
 try:
  if not {'Du_lieu','Thong_tin'}.issubset(wb.sheetnames):raise ValueError('Cần hai sheet Du_lieu và Thong_tin.')
  ws=wb['Du_lieu'];ms=wb['Thong_tin']
  for sheet in (ws,ms):
   if sheet.max_row is None or sheet.max_column is None:sheet.calculate_dimension(force=True)
  if ws.max_row>100 or ws.max_column>10 or ms.max_row>40 or ms.max_column>5:raise ValueError('Cấu trúc tệp quá lớn; dùng mẫu gốc.')
  rows=list(ws.values)
  if list(rows[0][:5]) != ['field','Chỉ tiêu','t','t1','t2']:raise ValueError('Không đổi tên 5 cột đầu trong Du_lieu.')
  d=blank();seen=set()
  for row in rows[1:]:
   key=row[0]
   if not key:continue
   if key not in FIELDS or key in seen:raise ValueError(f'field lạ hoặc trùng: {key}')
   seen.add(key)
   for i,y in enumerate(YEARS,2):
    value=row[i] if i<len(row) else None
    if value in (None,''):value=None
    elif isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):raise ValueError(f'{key}/{y}: nhập số Excel, không nhập công thức hoặc số dạng chữ.')
    d['periods'][y][key]=value
  if seen != set(FIELDS):raise ValueError('Thiếu dòng chỉ tiêu; dùng mẫu đầy đủ.')
  if list(next(ms.values)[:2])!=['key','value']:raise ValueError('Thong_tin cần cột key,value.')
  seen_meta=set()
  for row in list(ms.values)[1:]:
   key,value=row[:2]
   if not key:continue
   if key not in d['meta'] or key in seen_meta:raise ValueError(f'Thông tin lạ hoặc trùng: {key}')
   seen_meta.add(key)
   if key in BOOL_KEYS:
    if value is None:value=d['meta'][key]
    if not isinstance(value,bool):raise ValueError(f'{key}: dùng TRUE hoặc FALSE.')
   elif key=='year':
    if isinstance(value,bool) or not isinstance(value,(int,float)) or int(value)!=value or not 1990<=value<=2100:raise ValueError('Năm báo cáo phải là số nguyên 1990–2100.')
    value=int(value)
   else:
    if value is not None and not isinstance(value,str):raise ValueError(f'{key}: cần văn bản.')
    value=value or ''
    if value.startswith('='):raise ValueError('Không dùng công thức trong Thong_tin.')
   d['meta'][key]=value
  choices={'scope':['Hợp nhất','Riêng lẻ'],'unit':['VND','Nghìn VND','Triệu VND','Tỷ VND'],'sector':['Doanh nghiệp phi tài chính','Ngân hàng / Bảo hiểm / Chứng khoán'],'issuance':['Có','Không','Chưa xác định']}
  for k,values in choices.items():
   if d['meta'][k] not in values:raise ValueError(f'{k}: giá trị không hợp lệ.')
  if len(d['meta']['ticker'])>20 or len(d['meta']['source'])>2000:raise ValueError('Mã hoặc nguồn quá dài.')
  return d
 finally:wb.close()
